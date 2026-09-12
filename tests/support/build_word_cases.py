# -*- coding: utf-8 -*-
"""构造首批可重复审核的合成 Word 测试案例。"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt
from lxml import etree


##### 路径与版本板块 #####


REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
WORD_DIR = REPOSITORY_ROOT / "tests" / "fixtures" / "word"
BASELINE_DIR = REPOSITORY_ROOT / "tests" / "baselines" / "synthetic"
FIXTURE_VERSION = "1.1.0"
FIXED_TIME = datetime(2026, 8, 27, 0, 0, 0, tzinfo=timezone.utc)
ZIP_TIME = (2026, 8, 27, 0, 0, 0)


##### 通用构造板块 #####


def _prepare_document(title: str, subject: str) -> Document:
    document = Document()
    properties = document.core_properties
    properties.title = title
    properties.subject = subject
    properties.author = "Scholar Agent Synthetic Test"
    properties.keywords = "synthetic,word,latex,test"
    properties.comments = "完全合成测试文档，不包含真实论文内容。"
    properties.created = FIXED_TIME
    properties.modified = FIXED_TIME
    properties.last_printed = FIXED_TIME
    properties.revision = 1
    return document


def _normalize_docx_package(path: Path) -> None:
    """固定 DOCX ZIP 条目顺序与时间，减少环境无关的二进制漂移。"""
    temporary = path.with_suffix(".normalized.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9,
    ) as output:
        for name in sorted(source.namelist()):
            original = source.getinfo(name)
            info = zipfile.ZipInfo(name, ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = original.external_attr
            info.create_system = original.create_system
            output.writestr(info, source.read(name))
    os.replace(temporary, path)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _semantic_sha256(document: Document) -> str:
    """用段落与表格可见文字生成与 ZIP 元数据无关的语义哈希。"""
    records = [f"P\t{paragraph.style.name}\t{paragraph.text}" for paragraph in document.paragraphs]
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            records.append(
                f"T\t{table_index}\t{row_index}\t"
                + "\t".join(cell.text for cell in row.cells)
            )
    # python-docx 的 paragraph.text 不包含 OMML，必须额外纳入原生公式 XML。
    for formula in document.element.xpath(".//*[local-name()='oMathPara' or local-name()='oMath']"):
        parent = formula.getparent()
        if parent is not None and parent.tag.endswith(("oMathPara", "oMath")):
            continue
        records.append(f"M\t{etree.tostring(formula, encoding='unicode')}")
    return hashlib.sha256("\n".join(records).encode("utf-8")).hexdigest()


def _math_run(text: str) -> Any:
    run = OxmlElement("m:r")
    value = OxmlElement("m:t")
    value.text = text
    run.append(value)
    return run


def _append_inline_omml(paragraph: Any, text: str) -> None:
    formula = OxmlElement("m:oMath")
    formula.append(_math_run(text))
    paragraph._p.append(formula)


def _append_display_fraction_omml(paragraph: Any, numerator: str, denominator: str) -> None:
    """构造 Word 原生行间分式，验证结构化 OMML 而不只是纯文本节点。"""
    formula_paragraph = OxmlElement("m:oMathPara")
    formula = OxmlElement("m:oMath")
    fraction = OxmlElement("m:f")
    num = OxmlElement("m:num")
    num.append(_math_run(numerator))
    den = OxmlElement("m:den")
    den.append(_math_run(denominator))
    fraction.extend((num, den))
    formula.append(fraction)
    formula_paragraph.append(formula)
    paragraph._p.append(formula_paragraph)


def _save_case(
    case_id: str,
    filename: str,
    document: Document,
    expectation: dict[str, Any],
) -> None:
    WORD_DIR.mkdir(parents=True, exist_ok=True)
    BASELINE_DIR.mkdir(parents=True, exist_ok=True)
    path = WORD_DIR / filename
    semantic_hash = _semantic_sha256(document)
    document.save(path)
    _normalize_docx_package(path)
    expectation.update({
        "schema_version": "1.0.0",
        "case_id": case_id,
        "fixture_version": FIXTURE_VERSION,
        "source_kind": "fully_synthetic",
        "artifact": {
            "relative_path": path.relative_to(REPOSITORY_ROOT).as_posix(),
            "package_sha256": _sha256(path),
            "semantic_sha256": semantic_hash,
        },
    })
    target = BASELINE_DIR / f"{path.stem}.expected.json"
    target.write_text(
        json.dumps(expectation, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


##### W01 最小正文板块 #####


def _build_w01() -> tuple[Document, dict[str, Any]]:
    document = _prepare_document(
        "合成论文测试：最小正文", "验证最小正文、缺失元数据和零可选对象。",
    )
    document.add_paragraph("合成论文测试：最小正文", style="Title")
    document.add_heading("第一章 示例绪论", level=1)
    document.add_paragraph("本段用于验证普通中文正文能够完整进入模板。")
    document.add_paragraph("本段不包含图片、表格、公式和参考文献。")
    document.add_heading("1.1 研究目的", level=2)
    document.add_paragraph("本节用于验证二级标题和普通正文。")
    expectation = {
        "title": "最小正文与缺失元数据",
        "purpose": ["minimal_body", "missing_metadata", "zero_optional_objects"],
        "target_expectations": {
            "logical_inventory": {
                "title": 1, "level_1_headings": 1, "level_2_headings": 1,
                "body_paragraphs": 3, "images": 0, "tables": 0,
                "formulas": 0, "body_citations": 0, "reference_entries": 0,
            },
            "metadata": {
                "known": ["title"],
                "placeholders": ["author", "student_id", "advisor", "department"],
            },
            "recognition": {
                "requires_consolidated_review": True,
                "expected_high_confidence_headings": 2,
            },
            "references": {"mode": "none", "publish_policy": "not_applicable"},
            "dual_template": {
                "ouc_graduate": "degraded", "ouc_bachelor": "degraded",
                "allowed_degradations": ["missing_metadata"],
            },
        },
        "review_notes": [
            "标题、正文和层级必须保持，模板必需元数据使用待填写占位。",
            "零图片、零表格、零公式和零参考文献不得被当作失败。",
        ],
        "known_current_gaps": [
            "当前正式链路仍只有研究生模板适配器，本科双模板生成尚未实现。",
            "缺失元数据集中占位和零可选对象的双模板终态尚待实现验证。",
        ],
    }
    return document, expectation


##### W03 歧义结构板块 #####


def _format_noise_heading(paragraph, *, size: int, centered: bool = False) -> None:
    paragraph.style = "Normal"
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    )
    for run in paragraph.runs:
        run.bold = True
        run.font.size = Pt(size)


def _build_w03() -> tuple[Document, dict[str, Any]]:
    document = _prepare_document(
        "合成论文测试：结构歧义", "验证正文样式标题、第一表格误判和格式噪声。",
    )
    document.add_paragraph("合成论文测试：结构歧义", style="Title")
    chapter = document.add_paragraph("1 绪论")
    _format_noise_heading(chapter, size=16)
    document.add_paragraph("本章说明格式混乱的 Word 仍需保持正文顺序。")
    section = document.add_paragraph("1.1 研究背景")
    _format_noise_heading(section, size=14)
    document.add_paragraph("本段讨论2024年度研究结果，但它本身不是一个标题。")
    method = document.add_paragraph("研究方法")
    _format_noise_heading(method, size=14, centered=True)
    document.add_paragraph("该段落用于验证无编号标题需要集中确认。")

    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    values = (("监测时刻", "温度"), ("09:00", "21.5"), ("12:00", "23.0"))
    for row_index, row in enumerate(values):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value

    result = document.add_paragraph("实验结果")
    _format_noise_heading(result, size=14)
    result.runs[0].font.color.rgb = None
    document.add_paragraph("结果段落用于验证第二个无编号标题候选。")
    expectation = {
        "title": "标题样式混乱与普通首表",
        "purpose": ["ambiguous_headings", "format_noise", "first_table_not_cover"],
        "target_expectations": {
            "logical_inventory": {
                "title": 1, "numbered_heading_candidates": 2,
                "unnumbered_heading_candidates": 2, "body_paragraphs": 4,
                "images": 0, "tables": 1, "formulas": 0,
                "body_citations": 0, "reference_entries": 0,
            },
            "recognition": {
                "requires_consolidated_review": True,
                "expected_high_confidence_texts": ["1 绪论", "1.1 研究背景"],
                "expected_review_texts": ["研究方法", "实验结果"],
                "must_remain_body": ["本段讨论2024年度研究结果，但它本身不是一个标题。"],
                "first_table_role": "content_table",
            },
            "word_format_policy": "ignore_for_latex_output",
            "dual_template": {
                "ouc_graduate": "after_confirmation",
                "ouc_bachelor": "after_confirmation",
            },
        },
        "review_notes": [
            "字号、加粗、颜色和居中只作为弱证据，不作为最终 LaTeX 格式。",
            "研究方法和实验结果属于低置信度标题候选，用户确认前继续作为正文保存。",
            "第一张表是温度数据表，绝不能因位置被识别为封面信息。",
        ],
        "known_current_gaps": [
            "当前 recognize 结果能识别编号章，但尚未形成无编号标题的集中 Agent 建议。",
            "当前无题注普通表格虽被抽取，尚未作为可达正文对象进入模板渲染。",
        ],
    }
    return document, expectation


##### W05 完整引用板块 #####


def _build_w05() -> tuple[Document, dict[str, Any]]:
    document = _prepare_document(
        "合成论文测试：完整数字引用", "验证数字引用、文末条目和方括号非引用。",
    )
    document.add_paragraph("合成论文测试：完整数字引用", style="Title")
    document.add_heading("第一章 引用测试", level=1)
    document.add_paragraph("已有研究给出了相关结果[1]。")
    document.add_paragraph("另一项分析对该方法进行了扩展[2-3]。")
    document.add_paragraph("归一化区间为[0,1]，这里不是参考文献引用。")
    document.add_paragraph("数组索引示例为[1,2]，这里也不是参考文献引用。")
    for index in range(4):
        document.add_paragraph(f"合成填充正文{index + 1}，用于稳定文末区域位置。")
    document.add_heading("参考文献", level=1)
    entries = (
        "[1] Zhang A, Li B. Synthetic reference one[J]. Test Journal, 2024, 1(1): 1-10.",
        "[2] Wang C. Synthetic reference two[M]. Test Press, 2023.",
        "[3] Smith D. Synthetic reference three[C]. Test Conference, 2022: 20-25.",
    )
    for entry in entries:
        document.add_paragraph(entry, style="List Paragraph")
    expectation = {
        "title": "完整数字引用与方括号非引用",
        "purpose": ["word_numeric_references", "citation_context", "reference_fidelity"],
        "target_expectations": {
            "logical_inventory": {
                "title": 1, "level_1_headings": 1, "body_paragraphs": 8,
                "images": 0, "tables": 0, "formulas": 0,
                "body_citations": 2, "non_citation_brackets": 2,
                "reference_entries": 3,
            },
            "references": {
                "mode": "word_numeric_list",
                "body_citation_texts": ["[1]", "[2-3]"],
                "non_citation_texts": ["[0,1]", "[1,2]"],
                "number_to_key": {
                    "1": "word-ref-0001", "2": "word-ref-0002", "3": "word-ref-0003",
                },
                "explicit_labels_required": True,
                "preserve_reference_text": True,
            },
            "dual_template": {"ouc_graduate": "success", "ouc_bachelor": "success"},
        },
        "known_current_gaps": [
            "当前纯正则实现可能把语境中的[1,2]误识别为引用。",
            "当前抽取层可能把文末条目前缀[1]/[2]/[3]也创建为正文引用单元。",
        ],
        "review_notes": [
            "三个参考文献条目完全合成，不表示真实公开文献。",
            "引用转换必须保持原编号、原位置和文末条目可见文字。",
        ],
    }
    return document, expectation


##### W07 有效混合公式板块 #####


def _build_w07() -> tuple[Document, dict[str, Any]]:
    document = _prepare_document(
        "合成论文测试：有效混合公式", "验证 Word 原生 OMML 和 LaTeX 文本公式共存。",
    )
    document.add_paragraph("合成论文测试：有效混合公式", style="Title")
    document.add_heading("第一章 公式测试", level=1)

    inline_omml = document.add_paragraph("Word 原生行内公式")
    _append_inline_omml(inline_omml, "x+1")
    inline_omml.add_run("应保持在句内原位置。")

    document.add_paragraph("下式为 Word 原生行间分式：")
    display_omml = document.add_paragraph()
    _append_display_fraction_omml(display_omml, "a", "b")
    document.add_paragraph("其中，a 与 b 为合成变量。")

    document.add_paragraph("正文中的 LaTeX 行内公式 $a^2+b^2=c^2$ 应保持行内。")
    document.add_paragraph(r"$$\int_0^1 x^2\,dx=\frac{1}{3}$$")
    document.add_paragraph("归一化区间为[0,1]，不能识别为引用或公式。")
    document.add_paragraph(r"预算写作 \$100 时，转义美元符号不能开启公式。")

    expectation = {
        "title": "有效 OMML 与 LaTeX 文本公式混合",
        "purpose": [
            "omml_inline", "omml_display", "latex_inline", "latex_display",
            "formula_false_positive_guards",
        ],
        "terminology": {
            "word_native_formula": "OMML (Office Math Markup Language)",
            "user_term_normalization": "OML 按 Word 实际格式解释为 OMML",
        },
        "target_expectations": {
            "logical_inventory": {
                "title": 1, "level_1_headings": 1, "body_paragraphs": 8,
                "images": 0, "tables": 0, "formulas": 4,
                "omml_inline": 1, "omml_display": 1,
                "latex_inline": 1, "latex_display": 1,
                "body_citations": 0, "reference_entries": 0,
            },
            "formulas": {
                "preserve_source_order": True,
                "preserve_semantics": True,
                "expected_latex_bodies": ["x+1", r"\frac{a}{b}", "a^2+b^2=c^2", r"\int_0^1 x^2\,dx=\frac{1}{3}"],
                "must_remain_plain_text": ["[0,1]", r"\$100"],
            },
            "dual_template": {"ouc_graduate": "success", "ouc_bachelor": "success"},
        },
        "review_notes": [
            "四个公式必须按原顺序进入内容模型，不能把 OMML 和 LaTeX 文本重复计数。",
            "Word 原生行间公式与双美元行间公式可以使用模板支持的等价环境渲染，但不得改变公式含义。",
            "[0,1] 和转义金额均为反例，不得成为引用或公式。",
        ],
        "known_current_gaps": [
            "本科模板适配器尚未实现，因此双模板成功终态仍待后续验证。",
        ],
    }
    return document, expectation


##### W09 异常公式板块 #####


def _build_w09() -> tuple[Document, dict[str, Any]]:
    document = _prepare_document(
        "合成论文测试：异常公式", "验证公式错误不能被静默吞掉或作为成功发布。",
    )
    document.add_paragraph("合成论文测试：异常公式", style="Title")
    document.add_heading("第一章 公式错误测试", level=1)
    document.add_paragraph("未闭合的行内公式为 $x+y。")
    document.add_paragraph("单双美元混用的公式为 $$x+y$。")
    document.add_paragraph(r"花括号未闭合的公式为 $$\frac{a}{b$$。")
    document.add_paragraph("错误公式之后的正文必须继续保留。")
    expectation = {
        "title": "公式语法错误与发布阻断",
        "purpose": ["formula_delimiter_errors", "formula_syntax_error", "publish_gate"],
        "target_expectations": {
            "logical_inventory": {
                "title": 1, "level_1_headings": 1, "body_paragraphs": 4,
                "images": 0, "tables": 0, "valid_formulas": 0,
                "malformed_formula_candidates": 3,
            },
            "formulas": {
                "preserve_raw_text": True,
                "must_report_candidates": ["$x+y", "$$x+y$", r"$$\frac{a}{b$$"],
                "must_not_silently_repair": True,
                "must_not_drop_following_body": True,
            },
            "publish_policy": "blocked_before_success",
            "dual_template": {"ouc_graduate": "blocked", "ouc_bachelor": "blocked"},
        },
        "review_notes": [
            "该案例用于失败门禁，不要求生成可交付 PDF。",
            "若后续允许 Agent 建议修复，仍必须展示原文、候选修改和验证结果并获得授权。",
        ],
        "known_current_gaps": [
            "当前公式抽取正则可能忽略未闭合定界符，或把单双美元混用片段误识别为合法行内公式。",
            "公式语法预检与双模板统一阻断证据仍需在后续实现阶段验证。",
        ],
    }
    return document, expectation


##### 命令入口板块 #####


def build_all() -> list[Path]:
    cases = (
        ("W01", "W01_minimal_body.docx", _build_w01),
        ("W03", "W03_ambiguous_structure.docx", _build_w03),
        ("W05", "W05_references_complete.docx", _build_w05),
        ("W07", "W07_formulas_mixed_valid.docx", _build_w07),
        ("W09", "W09_formulas_malformed.docx", _build_w09),
    )
    outputs = []
    for case_id, filename, builder in cases:
        document, expectation = builder()
        _save_case(case_id, filename, document, expectation)
        outputs.append(WORD_DIR / filename)
    return outputs


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="连续生成两次并检查文件哈希稳定。")
    arguments = parser.parse_args()
    first = build_all()
    first_hashes = {path.name: _sha256(path) for path in first}
    if arguments.check:
        second = build_all()
        second_hashes = {path.name: _sha256(path) for path in second}
        if first_hashes != second_hashes:
            raise RuntimeError(f"合成 Word 二次生成哈希不稳定：{first_hashes} != {second_hashes}")
    print(json.dumps(first_hashes, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
