# -*- coding: utf-8 -*-
"""G8-A 上传、持久任务、结构预览与确认集成测试。"""

from __future__ import annotations

import io
import json
import os
import time
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
import pytest

from infrastructure import stage1_adapter
from infrastructure.stage1_adapter import WorkspaceUnavailableError
from main import app
from persistence.store import ProjectStore, get_store
from infrastructure.tasks import TaskManager, get_task_manager
from generation.service import execute_generation
from tests.support.delivery import delivery_values, frozen_values


##### 测试夹具板块 #####


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("第一章 绪论", level=1)
    document.add_paragraph("这是用于接口闭环测试的正文。")
    document.add_heading("1.1 研究背景", level=2)
    document.add_paragraph("正文中包含公式 $a+b=c$，但接口不会改写它。")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _w03_bytes() -> bytes:
    return (
        Path(__file__).resolve().parents[3]
        / "tests/fixtures/word/W03_ambiguous_structure.docx"
    ).read_bytes()


def _w09_bytes() -> bytes:
    return (
        Path(__file__).resolve().parents[3]
        / "tests/fixtures/word/W09_formulas_malformed.docx"
    ).read_bytes()


def _client(tmp_path: Path) -> tuple[TestClient, ProjectStore, TaskManager]:
    workspace = tmp_path / "workspace"
    database = tmp_path / "scholar.db"
    os.environ["SCHOLAR_WORKSPACE_ROOT"] = str(workspace)
    stage1_adapter._get_workspace_manager.cache_clear()
    store = ProjectStore(database)
    tasks = TaskManager(database, max_workers=1)
    app.dependency_overrides[get_store] = lambda: store
    app.dependency_overrides[get_task_manager] = lambda: tasks
    return TestClient(app), store, tasks


##### 纵向闭环板块 #####


def test_upload_extract_poll_preview_and_confirm(tmp_path: Path) -> None:
    client, store, _ = _client(tmp_path)
    try:
        created = client.post(
            "/api/projects",
            files={"docx": ("paper.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"template_id": "ouc-graduate", "reference_source": "auto"},
        )
        assert created.status_code == 200, created.text
        project_id = created.json()["project_id"]
        project = store.get(project_id)
        assert project is not None
        assert Path(project.docx_path).parent.name == "inputs"
        assert not hasattr(project, "img_dir")

        triggered = client.post(f"/api/projects/{project_id}/extract")
        assert triggered.status_code == 200
        task_id = triggered.json()["task_id"]
        for _ in range(100):
            task = client.get(f"/api/tasks/{task_id}").json()
            if task["status"] in {"done", "failed"}:
                break
            time.sleep(0.05)
        assert task["status"] == "done", task
        extraction = (
            Path(project.workspace_dir) / "reports" / "extraction" / "content_units.json"
        )
        assert extraction.is_file()

        with patch.object(
            stage1_adapter._get_runner(), "extract",
            side_effect=AssertionError("重复提取不得再次解析 Word"),
        ), patch.object(
            stage1_adapter._get_runner(), "recognize",
            side_effect=AssertionError("已持久化结构不得再次识别"),
        ):
            repeated_task_id = client.post(
                f"/api/projects/{project_id}/extract"
            ).json()["task_id"]
            for _ in range(100):
                repeated = client.get(f"/api/tasks/{repeated_task_id}").json()
                if repeated["status"] in {"done", "failed"}:
                    break
                time.sleep(0.05)
        assert repeated["status"] == "done", repeated

        preview = client.get(f"/api/projects/{project_id}/structure")
        assert preview.status_code == 200, preview.text
        body = preview.json()
        assert body["source_sha256"] == created.json()["source_sha256"]
        assert body["schema_version"] == "1.7.0"
        assert body["headings"]

        conflict = client.put(
            f"/api/projects/{project_id}/structure",
            json={
                "base_revision": 99,
                "heading_overrides": [],
                "object_binding_overrides": [],
                "confirmed": True,
            },
        )
        assert conflict.status_code == 409
        assert conflict.json()["error"] == "revision_conflict"

        first = body["headings"][0]
        confirmed = client.put(
            f"/api/projects/{project_id}/structure",
            json={
                "base_revision": body["revision"],
                "heading_overrides": [{"unit_id": first["unit_id"], "level": "chapter"}],
                "object_binding_overrides": [],
                "confirmed": True,
            },
        )
        assert confirmed.status_code == 200, confirmed.text
        assert confirmed.json()["revision"] == 1
        assert confirmed.json()["confirmed"] is True
        assert client.get(f"/api/projects/{project_id}/structure").json()["revision"] == 1

        # 内容模型升级后的新识别快照必须使旧确认失效，不能继续返回旧版本。
        project = store.get(project_id)
        assert project is not None
        recognition = Path(project.workspace_dir) / "reports" / "recognition.json"
        upgraded = json.loads(recognition.read_text(encoding="utf-8"))
        upgraded["schema_version"] = "1.7.0-test"
        upgraded["revision"] = 0
        upgraded["confirmed"] = False
        recognition.write_text(
            json.dumps(upgraded, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        refreshed = client.get(f"/api/projects/{project_id}/structure").json()
        assert refreshed["schema_version"] == "1.7.0-test"
        assert refreshed["revision"] == 0
        assert refreshed["confirmed"] is False
    finally:
        app.dependency_overrides.clear()
        stage1_adapter._get_workspace_manager.cache_clear()


def test_ambiguous_headings_must_be_resolved_in_one_confirmation(tmp_path: Path) -> None:
    client, _, _ = _client(tmp_path)
    try:
        created = client.post(
            "/api/projects",
            files={"docx": ("w03.docx", _w03_bytes())},
            data={"template_id": "ouc-graduate", "reference_source": "auto"},
        )
        project_id = created.json()["project_id"]
        task_id = client.post(f"/api/projects/{project_id}/extract").json()["task_id"]
        for _ in range(100):
            task = client.get(f"/api/tasks/{task_id}").json()
            if task["status"] in {"done", "failed"}:
                break
            time.sleep(0.05)
        assert task["status"] == "done", task
        snapshot = client.get(f"/api/projects/{project_id}/structure").json()
        review = [item for item in snapshot["headings"] if item["requires_review"]]
        assert [item["text"] for item in review] == ["研究方法", "实验结果"]

        rejected = client.put(
            f"/api/projects/{project_id}/structure",
            json={"base_revision": 0, "heading_overrides": [],
                  "object_binding_overrides": [], "confirmed": True},
        )
        assert rejected.status_code == 422
        assert rejected.json()["error"] == "unresolved_structure_review"

        accepted = client.put(
            f"/api/projects/{project_id}/structure",
            json={
                "base_revision": 0,
                "heading_overrides": [
                    {"unit_id": item["unit_id"], "level": "body"}
                    for item in review
                ],
                "object_binding_overrides": [], "confirmed": True,
            },
        )
        assert accepted.status_code == 200, accepted.text
        assert all(
            item["review_status"] == "resolved"
            for item in accepted.json()["headings"] if item["requires_review"]
        )
    finally:
        app.dependency_overrides.clear()
        stage1_adapter._get_workspace_manager.cache_clear()


def test_malformed_formulas_fail_before_compile_and_leave_only_report(
    tmp_path: Path,
) -> None:
    client, _, _ = _client(tmp_path)
    try:
        created = client.post(
            "/api/projects", files={"docx": ("w09.docx", _w09_bytes())}
        )
        project_id = created.json()["project_id"]
        task_id = client.post(f"/api/projects/{project_id}/extract").json()["task_id"]
        for _ in range(100):
            task = client.get(f"/api/tasks/{task_id}").json()
            if task["status"] in {"done", "failed"}:
                break
            time.sleep(0.05)
        assert task["status"] == "done", task

        snapshot = client.get(f"/api/projects/{project_id}/structure").json()
        assert snapshot["formula_review"]["malformed_formula_count"] == 3
        confirmed = client.put(
            f"/api/projects/{project_id}/structure",
            json={"base_revision": 0, "heading_overrides": [],
                  "object_binding_overrides": [], "confirmed": True},
        )
        assert confirmed.status_code == 200, confirmed.text

        with patch(
            "pipeline_api.PipelineRunner.extract",
            side_effect=AssertionError("生成阶段不得再次解析 Word"),
        ), patch(
            "pipeline_api.PipelineRunner.recognize",
            side_effect=AssertionError("生成阶段不得再次识别结构"),
        ):
            triggered = client.post(
                f"/api/projects/{project_id}/generate",
                json={"structure_revision": 1},
            )
            assert triggered.status_code == 200, triggered.text
            generation_id = triggered.json()["generation_id"]
            generation_task_id = triggered.json()["task_id"]
            for _ in range(100):
                task = client.get(f"/api/tasks/{generation_task_id}").json()
                if task["status"] in {"done", "failed"}:
                    break
                time.sleep(0.05)
        assert task["status"] == "done", task

        generation = client.get(f"/api/generations/{generation_id}").json()
        assert generation["status"] == "failed"
        assert generation["quality_status"] == "blocked"
        assert not any(item["kind"] in {"pdf", "latex_source"} for item in generation["artifacts"])
        report = client.get(
            f"/api/generations/{generation_id}/reports/pipeline_log.json"
        )
        assert report.status_code == 200
        payload = report.json()
        assert payload["formula_preflight"]["malformed_formula_count"] == 3
        assert payload["compile"] == {"success": False, "status": "not_run"}
        assert payload["llm_tasks"] == []
    finally:
        app.dependency_overrides.clear()
        stage1_adapter._get_workspace_manager.cache_clear()


##### 输入拒绝板块 #####


def test_object_confirmation_rejects_unresolved_duplicate_and_stale_schema(tmp_path):
    document = Document()
    document.add_heading("第一章 测试", 1)
    document.add_paragraph("表1-1 第一组数据")
    for number in range(3):
        if number == 2:
            document.add_paragraph("表1-2 第二组数据")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "指标"
        table.cell(0, 1).text = "数值"
        table.cell(1, 0).text = "对象"
        table.cell(1, 1).text = str(number)
    document.add_paragraph("正文结束。")
    buffer = io.BytesIO()
    document.save(buffer)
    client, store, _ = _client(tmp_path)
    try:
        created = client.post("/api/projects", files={"docx": ("objects.docx", buffer.getvalue())})
        project_id = created.json()["project_id"]
        task_id = client.post(f"/api/projects/{project_id}/extract").json()["task_id"]
        for _ in range(100):
            task = client.get(f"/api/tasks/{task_id}").json()
            if task["status"] in {"done", "failed"}:
                break
            time.sleep(0.05)
        assert task["status"] == "done", task
        url = f"/api/projects/{project_id}/structure"
        snapshot = client.get(url).json()
        first, second = snapshot["object_bindings"]
        assert first["status"] == "needs_review" and len(first["candidates"]) == 2
        assert client.put(url, json={"base_revision": 0}).status_code == 422
        duplicate = client.put(url, json={"base_revision": 0, "object_binding_overrides": [
            {"caption_unit_id": first["caption_unit_id"], "object_unit_ids": second["object_unit_ids"]},
        ]})
        assert duplicate.status_code == 422 and "重复绑定" in duplicate.text
        accepted = client.put(url, json={"base_revision": 0, "object_binding_overrides": [
            {"caption_unit_id": first["caption_unit_id"], "object_unit_ids": first["candidates"][0]["object_unit_ids"]},
        ]})
        assert accepted.status_code == 200, accepted.text
        assert accepted.json()["counts"]["object_binding_reviews"] == 0
        project = store.get(project_id)
        path = Path(project.workspace_dir) / "reports/confirmed_structure.json"
        saved = json.loads(path.read_text(encoding="utf-8"))
        assert saved["chapters"] and "source_sha256" not in saved
        saved["schema_version"] = "1.6.0"
        path.write_text(json.dumps(saved), encoding="utf-8")
        recognition = Path(project.workspace_dir) / "reports/recognition.json"
        raw = json.loads(recognition.read_text(encoding="utf-8"))
        raw["schema_version"] = "1.6.0"
        recognition.write_text(json.dumps(raw), encoding="utf-8")
        assert client.put(url, json={"base_revision": 1}).status_code == 422
    finally:
        app.dependency_overrides.clear()
        stage1_adapter._get_workspace_manager.cache_clear()


def test_citation_confirmation_rejects_unknown_unresolved_and_stale_revision(tmp_path):
    from tests.unit.content_extraction.test_citations import citation_document

    document = citation_document()
    buffer = io.BytesIO()
    document.save(buffer)
    client, _, _ = _client(tmp_path)
    try:
        created = client.post("/api/projects", files={"docx": ("citations.docx", buffer.getvalue())}, data={"reference_source": "word_list"})
        project_id = created.json()["project_id"]
        task_id = client.post(f"/api/projects/{project_id}/extract").json()["task_id"]
        for _ in range(100):
            task = client.get(f"/api/tasks/{task_id}").json()
            if task["status"] in {"done", "failed"}:
                break
            time.sleep(0.05)
        assert task["status"] == "done", task
        url = f"/api/projects/{project_id}/structure"
        snapshot = client.get(url).json()
        pending = [item for item in snapshot["citation_review"]["decisions"] if item["decision"] == "needs_review"]
        assert len(pending) == 1
        assert client.put(url, json={"base_revision": 0}).status_code == 422
        assert client.post(f"/api/projects/{project_id}/generate", json={"structure_revision": 1}).status_code == 409
        for override in [
            {"unit_id": "unknown", "decision": "body_citation"},
            {"unit_id": pending[0]["unit_id"], "decision": "reference_label"},
            {"unit_id": pending[0]["unit_id"], "decision": "body_citation", "numbers": ["2"]},
        ]:
            assert client.put(url, json={"base_revision": 0, "citation_overrides": [override]}).status_code == 422
        override = {"unit_id": pending[0]["unit_id"], "decision": "non_citation"}
        accepted = client.put(url, json={"base_revision": 0, "citation_overrides": [override]})
        assert accepted.status_code == 200, accepted.text
        assert accepted.json()["citation_review"]["counts"]["needs_review"] == 0
        assert client.put(url, json={"base_revision": 0, "citation_overrides": [override]}).status_code == 409
    finally:
        app.dependency_overrides.clear()
        stage1_adapter._get_workspace_manager.cache_clear()


def test_rejects_fake_docx_and_client_paths(tmp_path: Path) -> None:
    client, _, _ = _client(tmp_path)
    try:
        response = client.post(
            "/api/projects",
            files={"docx": ("paper.docx", b"not-a-zip", "application/octet-stream")},
            data={"docx_path": "C:/danger.docx", "template_id": "ouc-graduate"},
        )
        assert response.status_code == 422
        assert response.json()["error"] == "upload_invalid"
    finally:
        app.dependency_overrides.clear()
        stage1_adapter._get_workspace_manager.cache_clear()


def test_workspace_failure_is_readable_cross_origin_response(tmp_path: Path) -> None:
    client, _, _ = _client(tmp_path)
    try:
        with patch(
            "api.routes.projects.create_project_workspace",
            side_effect=WorkspaceUnavailableError("项目工作区不可用，请检查服务端运行目录配置。"),
        ):
            response = client.post(
                "/api/projects",
                files={"docx": ("paper.docx", _docx_bytes())},
                headers={"Origin": "http://127.0.0.1:5173"},
            )
        assert response.status_code == 503
        assert response.json() == {
            "error": "workspace_unavailable",
            "detail": "项目工作区不可用，请检查服务端运行目录配置。",
        }
        assert response.headers["access-control-allow-origin"] == "*"
    finally:
        app.dependency_overrides.clear()
        stage1_adapter._get_workspace_manager.cache_clear()


##### generation 接口板块 #####


def test_generate_requires_confirmed_current_revision(tmp_path: Path) -> None:
    client, store, _ = _client(tmp_path)
    try:
        created = client.post(
            "/api/projects", files={"docx": ("paper.docx", _docx_bytes())}
        )
        project_id = created.json()["project_id"]
        rejected = client.post(
            f"/api/projects/{project_id}/generate",
            json={"structure_revision": 1},
        )
        assert rejected.status_code == 409
        assert rejected.json()["error"] == "structure_not_confirmed"
        assert store.get(project_id) is not None
    finally:
        app.dependency_overrides.clear()
        stage1_adapter._get_workspace_manager.cache_clear()


def test_artifact_path_escape_is_rejected(tmp_path: Path) -> None:
    client, store, _ = _client(tmp_path)
    try:
        store.create(project_id="p1", docx_path="paper.docx", bib_path=None, citation_map_path=None,
            template_id="ouc-graduate", reference_source="word_list", workspace_dir=str(tmp_path / "project"),
            source_sha256="hash")
        outside = tmp_path / "secret.pdf"
        outside.write_bytes(b"%PDF-secret")
        generation = store.create_generation(
            generation_id="g-path-test", project_id="p1",
            structure_revision=1, source_sha256="hash",
        )
        values = delivery_values(store, generation.generation_id)
        store.publish_trusted_generation(generation.generation_id, **values)
        values["artifact_manifest"]["artifacts"][0]["relative_path"] = "../secret.pdf"
        with store._connect() as connection:
            connection.execute("UPDATE generations SET artifact_manifest_json=? WHERE generation_id=?",
                (json.dumps(values["artifact_manifest"]), generation.generation_id))
        response = client.get(f"/api/generations/{generation.generation_id}/pdf")
        assert response.status_code == 404
        assert response.json()["error"] == "artifact_not_found"
    finally:
        app.dependency_overrides.clear()
        stage1_adapter._get_workspace_manager.cache_clear()


def test_pdf_artifact_is_served_inline_for_browser_preview(tmp_path: Path) -> None:
    client, store, _ = _client(tmp_path)
    try:
        store.create(project_id="p1", docx_path="paper.docx", bib_path=None, citation_map_path=None,
            template_id="ouc-graduate", reference_source="word_list", workspace_dir=str(tmp_path / "project"),
            source_sha256="hash")
        generation = store.create_generation(
            generation_id="g-preview", project_id="p1",
            structure_revision=1, source_sha256="hash",
        )
        store.publish_trusted_generation(generation.generation_id, **delivery_values(store, generation.generation_id))

        response = client.get(f"/api/generations/{generation.generation_id}/pdf")

        assert response.status_code == 200
        assert response.headers["content-type"] == "application/pdf"
        assert response.headers["content-disposition"].startswith("inline;")
    finally:
        app.dependency_overrides.clear()
        stage1_adapter._get_workspace_manager.cache_clear()


def test_generation_internal_error_is_persisted(tmp_path: Path) -> None:
    database = tmp_path / "state.db"
    store = ProjectStore(database)
    store.create(
        project_id="p1", docx_path="paper.docx", bib_path=None,
        citation_map_path=None, template_id="ouc-graduate", reference_source="auto",
        workspace_dir=str(tmp_path / "project"), source_sha256="hash", structure_revision=1,
    )
    store.create_generation(
        generation_id="g-failed", project_id="p1",
        structure_revision=1, source_sha256="hash",
        **frozen_values(store.get("p1"), "g-failed"),
    )
    snapshot = tmp_path / "project" / "reports" / "confirmed_structure.json"
    snapshot.parent.mkdir(parents=True, exist_ok=True)
    snapshot.write_text("{}", encoding="utf-8")
    failure_report = snapshot.parent / "g-failed" / "pipeline_log.json"
    failure_report.parent.mkdir()
    failure_report.write_text('{"published": false}', encoding="utf-8")
    with patch(
        "generation.service.run_project_generation",
        side_effect=RuntimeError("injected failure"),
    ), pytest.raises(RuntimeError, match="injected failure"):
        execute_generation(
            database_path=str(database), generation_id="g-failed",
        )
    restored = store.get_generation("g-failed")
    assert restored is not None
    assert restored.status == "failed"
    assert restored.quality_status == "internal_error"
    assert any(item["name"] == "pipeline_log.json" for item in restored.artifact_manifest["artifacts"])
    with patch('generation.service.run_project_generation', return_value={
        'status': 'failed', 'quality_status': 'internal_error', 'output_dir': None,
        'report_dir': str(failure_report.parent), 'artifact_manifest': {'artifacts': []},
        'detail': '门禁内部异常', 'gate_snapshot': {'published': False, 'gate_statuses': {}},
    }):
        execute_generation(database_path=str(database), generation_id='g-failed')
    assert store.get_generation('g-failed').quality_status == 'internal_error'
