# -*- coding: utf-8 -*-
"""通用 Word 抽取与识别的单元、虚拟文档和真实论文节选测试。"""

from __future__ import annotations

import base64
import tempfile
import unittest
from pathlib import Path

import pytest

from docx import Document
from docx.oxml import OxmlElement

from tests.support.paths import GRADUATE_TEMPLATE


##### 测试环境板块 #####


from pipeline.structural_gate import structural_fidelity_gate
from pipeline_api import PipelineRunner


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUB"
    "AScY42YAAAAASUVORK5CYII="
)


def _append_simple_omml(paragraph, text: str = "x+1") -> None:
    formula = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    value = OxmlElement("m:t")
    value.text = text
    run.append(value)
    formula.append(run)
    paragraph._p.append(formula)


##### 虚拟 Word 测试板块 #####


class OrderedExtractionTests(unittest.TestCase):
    def test_order_formula_image_and_table_binding(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            image_path = root / "pixel.png"
            image_path.write_bytes(PNG_1X1)
            docx_path = root / "virtual-thesis.docx"

            document = Document()
            document.add_paragraph("表1-1 图表清单中的重复项")
            for index in range(6):
                document.add_paragraph(f"前置内容{index}")
            document.add_heading("第一章 绪论", level=1)
            document.add_paragraph("正文第一段")
            document.add_picture(str(image_path))
            document.add_paragraph("图1-1 虚拟图片")
            document.add_paragraph("表1-1 虚拟数据表")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "指标"
            table.cell(0, 1).text = "数值"
            table.cell(1, 0).text = "A"
            table.cell(1, 1).text = "1"
            equation_paragraph = document.add_paragraph("公式：")
            _append_simple_omml(equation_paragraph)
            document.save(docx_path)

            runner = PipelineRunner()
            extracted = runner.extract(str(docx_path))
            recognized = runner.recognize(extracted)

            unit_types = [unit["unit_type"] for unit in extracted.content_units]
            self.assertIn("image", unit_types)
            self.assertIn("formula", unit_types)
            self.assertIn("table", unit_types)
            self.assertEqual(recognized.review["chapters"][0]["title"], "绪论")
            self.assertEqual(sum(item["kind"] == "table" and item["status"] == "bound" for item in recognized.review["object_bindings"]), 1)
            self.assertEqual(recognized.internal, {})

            formulas = [
                unit for unit in extracted.content_units if unit["unit_type"] == "formula"
            ]
            self.assertEqual(formulas[0]["text"], "x+1")

            table_unit = next(
                unit for unit in extracted.content_units if unit["unit_type"] == "table"
            )
            before = next(
                unit
                for unit in extracted.content_units
                if unit["unit_type"] == "paragraph" and unit["text"] == "表1-1 虚拟数据表"
            )
            self.assertLess(before["order"], table_unit["order"])

            bindings = recognized.review["object_bindings"]
            duplicate_statuses = [
                item["status"] for item in bindings if item["number"] == "表1-1"
            ]
            self.assertCountEqual(duplicate_statuses, ["bound"])

    def test_structural_count_has_no_sample_paragraph_window(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            image = root / "pixel.png"
            image.write_bytes(PNG_1X1)
            document = Document()
            document.add_heading("第一章 绪论", level=1)
            for index in range(38):
                document.add_paragraph(f"正文段落{index}。")
            document.add_picture(str(image))
            document.add_paragraph("图1-1 第四十段的真实图片")
            source = root / "input.docx"
            document.save(source)
            runner = PipelineRunner()
            extracted = runner.extract(str(source))
            recognized = runner.recognize(extracted)
            rendered = runner.render(extracted, recognized, str(GRADUATE_TEMPLATE), str(root / "output"), None)
            checks = structural_fidelity_gate(
                extracted.word_structure, root / "output" / runner.template_adapter.content_directory,
                confirmed_structure=recognized.review, main_tex_path=root / "output" / "main.tex",
                chapter_files=rendered.chapter_files, template_adapter=runner.template_adapter,
            )
            caption = recognized.review["object_bindings"][0]["caption_unit_id"]
            self.assertEqual(checks["figure_count"], (1, 1))
            self.assertEqual(checks[f"caption:{caption}"], (1, 1))


##### 尾部对象保真板块 #####


@pytest.mark.parametrize("trailing_blank", [False, True])
def test_trailing_data_table_survives_blank_paragraphs(tmp_path, trailing_blank):
    from pipeline.confirmed_structure import resolve_render_structure
    from pipeline.content_fidelity_gate import build_content_fidelity_report

    document = Document()
    document.add_heading("第一章 测试", level=1)
    document.add_paragraph("正文。")
    document.add_paragraph("")
    table = document.add_table(rows=2, cols=2)
    for cell, value in zip((cell for row in table.rows for cell in row.cells), ("项目", "数值", "样本", "7")):
        cell.text = value
    if trailing_blank:
        document.add_paragraph("")
    source = tmp_path / "trailing-table.docx"
    document.save(source)
    runner = PipelineRunner()
    extracted = runner.extract(str(source))
    recognized = resolve_render_structure(extracted, runner.recognize(extracted))
    rendered = runner.render(extracted, recognized, str(GRADUATE_TEMPLATE), str(tmp_path / "rendered"), None)
    table_id = next(unit["unit_id"] for unit in extracted.content_units if unit["unit_type"] == "table")
    assert any(item["marker_unit_id"] == table_id and item["role"] == "content_table" and item["status"] == "rendered" for item in rendered.render_trace["records"])
    assert build_content_fidelity_report(extracted, recognized, rendered)["all_passed"]


if __name__ == "__main__":
    unittest.main(verbosity=2)
