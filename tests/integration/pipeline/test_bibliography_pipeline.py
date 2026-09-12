# -*- coding: utf-8 -*-
"""虚拟 Word 参考文献降级、渲染与保真集成测试。"""

from pathlib import Path

from docx import Document

from pipeline.content_fidelity_gate import build_content_fidelity_report
from pipeline.structural_gate import (
    structural_fidelity_gate,
)
from pipeline_api import PipelineRunner
from tests.support.paths import GRADUATE_TEMPLATE


##### Word参考文献链路板块 #####


def test_word_only_references_render_and_pass_fidelity(tmp_path: Path) -> None:
    docx = tmp_path / "word-references.docx"
    document = Document()
    document.add_heading("第一章 绪论", level=1)
    document.add_paragraph("已有研究[1-2]给出了相关结果。")
    for index in range(5):
        document.add_paragraph(f"补充正文{index}。")
    document.add_heading("参考文献", level=1)
    document.add_paragraph("张三. 示例研究[J]. 示例期刊, 2023.", style="List Paragraph")
    document.add_paragraph("Li A. Another study[J]. Journal, 2024.", style="List Paragraph")
    document.save(docx)

    runner = PipelineRunner()
    extract = runner.extract(str(docx))
    recognize = runner.recognize(extract)
    output = tmp_path / "output"
    render = runner.render(
        extract,
        recognize,
        str(GRADUATE_TEMPLATE),
        str(output),
        None,
        reference_source="word",
    )
    report = build_content_fidelity_report(extract, recognize, render)
    structural = structural_fidelity_gate(
        extract.word_structure,
        output / "contents",
        main_tex_path=output / "main.tex",
        chapter_files=render.chapter_files,
        confirmed_structure=recognize.review,
        template_adapter=runner.template_adapter,
    )
    main = (output / "main.tex").read_text(encoding="utf-8")
    references = (output / "data" / "references.tex").read_text(encoding="utf-8")
    section = (output / render.chapter_files[0]).read_text(encoding="utf-8")

    assert report["all_passed"], report["blocking_or_unimplemented"]
    assert structural["chapter_count"] == (1, 1)
    assert len(render.chapter_files) == 1
    assert r"\input{data/references.tex}" in main
    assert r"\bibitem[1]{word-ref-0001}" in references
    assert r"\cite{word-ref-0001,word-ref-0002}" in section
    assert "参考文献" not in section
