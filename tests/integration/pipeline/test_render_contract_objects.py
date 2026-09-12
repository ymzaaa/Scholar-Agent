# -*- coding: utf-8 -*-
"""双 OUC 图表、双语题注、特殊文字和实际编译验收。"""

from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document

from pipeline.content_fidelity_gate import build_content_fidelity_report
from pipeline.structural_gate import structural_fidelity_gate
from pipeline_api import PipelineRunner
from template_registry.registry import resolve_builtin_template, load_builtin_template_registry
from tests.integration.pipeline.test_g4_word_objects import PNG_1X1
from tests.support.template_source import source_snapshot


##### 双语图表夹具板块 #####


def _document(root):
    image = root / 'pixel.png'
    image.write_bytes(PNG_1X1)
    document = Document()
    document.add_heading('第一章 测试_标题&字符', level=1)
    document.add_paragraph('正文 # % & _ ^ ~ { }，普通金额 \\$5。')
    document.add_paragraph('网址 https://example.org/a_b?q=one&percent=20%25#part_1。')
    document.add_picture(str(image))
    document.add_paragraph('图1-1 图片_示例')
    document.add_paragraph('Fig. 1-1 Image example')
    document.add_paragraph('表1-1 表格_示例')
    document.add_paragraph('Table. 1-1 Table example')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '第一行\n第二行'
    table.cell(0, 1).text = '数值'
    table.cell(1, 0).text = '$x_1$'
    table.cell(1, 1).text = '20%'
    path = root / 'objects.docx'
    document.save(path)
    return path


##### 对象与编译验收板块 #####


@pytest.mark.parametrize('template_id', ['ouc-graduate', 'ouc-bachelor'])
def test_bilingual_objects_and_plain_characters_compile_without_source_changes(tmp_path, template_id):
    resolved = resolve_builtin_template(template_id)
    runner = PipelineRunner(resolved_template=resolved)
    source = load_builtin_template_registry().source_root(resolved.manifest)
    before = source_snapshot(source)
    extracted = runner.extract(str(_document(tmp_path)))
    original = deepcopy(extracted.content_units)
    recognized = runner.recognize(extracted)
    output = tmp_path / 'output'
    rendered = runner.render(extracted, recognized, str(source), str(output), None)
    fidelity = build_content_fidelity_report(extracted, recognized, rendered)
    assert fidelity['all_passed'], fidelity['blocking_or_unimplemented']
    structural_fidelity_gate(extracted.word_structure, output / runner.template_adapter.content_directory,
        confirmed_structure=recognized.review, main_tex_path=output / runner.template_adapter.entrypoint,
        chapter_files=rendered.chapter_files, template_adapter=runner.template_adapter)
    for binding in recognized.review['object_bindings']:
        assert binding['caption_en_unit_id']
        record = next(item for item in rendered.render_trace['records'] if item['marker_unit_id'] == binding['caption_unit_id'])
        assert binding['caption_en_unit_id'] in record['source_unit_ids']
    compiled = runner.compile(str(output))
    assert compiled.success, (compiled.raw_error, compiled.process_results)
    assert extracted.content_units == original
    assert source_snapshot(source) == before
