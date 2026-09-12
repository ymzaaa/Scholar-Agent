# -*- coding: utf-8 -*-
"""G3-B2脚注、域、交叉引用、链接和文本框专项测试。"""

from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.opc.constants import RELATIONSHIP_TYPE
from lxml import etree

from tests.support.paths import GRADUATE_TEMPLATE


##### 测试环境板块 #####


from pipeline.content_fidelity_gate import build_content_fidelity_report
from pipeline_api import PipelineRunner


CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
RELATIONSHIPS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
FOOTNOTES_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
FOOTNOTES_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"


def _run_with_text(text: str) -> OxmlElement:
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    return run


def _append_bookmark(paragraph, name: str, bookmark_id: str = "42") -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bookmark_id)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _append_complex_field(paragraph, instruction: str, result: str) -> None:
    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)
    instruction_run = OxmlElement("w:r")
    instruction_node = OxmlElement("w:instrText")
    instruction_node.set(qn("xml:space"), "preserve")
    instruction_node.text = instruction
    instruction_run.append(instruction_node)
    separate_run = OxmlElement("w:r")
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run.append(separate)
    end_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    paragraph._p.extend([
        begin_run, instruction_run, separate_run, _run_with_text(result), end_run
    ])


def _append_simple_field(paragraph, instruction: str, result: str) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    field.append(_run_with_text(result))
    paragraph._p.append(field)


def _append_malformed_field(paragraph, instruction: str, result: str) -> None:
    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)
    instruction_run = OxmlElement("w:r")
    instruction_node = OxmlElement("w:instrText")
    instruction_node.text = instruction
    instruction_run.append(instruction_node)
    separate_run = OxmlElement("w:r")
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run.append(separate)
    paragraph._p.extend([
        begin_run, instruction_run, separate_run, _run_with_text(result)
    ])


def _append_hyperlink(paragraph, display: str, target: str) -> None:
    relationship_id = paragraph.part.relate_to(
        target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    hyperlink.append(_run_with_text(display))
    paragraph._p.append(hyperlink)


def _append_text_box(paragraph, text: str, *, nested_table: bool = False) -> None:
    table_xml = ""
    if nested_table:
        table_xml = (
            "<w:tbl><w:tr><w:tc><w:p><w:r><w:t>表格</w:t></w:r></w:p>"
            "</w:tc></w:tr></w:tbl>"
        )
    xml = (
        '<w:pict '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<v:shape id="TextBox1"><v:textbox><w:txbxContent>'
        f'<w:p><w:r><w:t>{text}</w:t></w:r>'
        '<m:oMath><m:r><m:t>y+1</m:t></m:r></m:oMath>'
        '<w:fldSimple w:instr=" HYPERLINK &quot;https://example.org&quot; ">'
        '<w:r><w:t>框内链接</w:t></w:r></w:fldSimple></w:p>'
        f'{table_xml}'
        '</w:txbxContent></v:textbox></v:shape></w:pict>'
    )
    paragraph._p.append(parse_xml(xml))


def _append_footnote_reference(paragraph, note_id: str) -> None:
    run = OxmlElement("w:r")
    reference = OxmlElement("w:footnoteReference")
    reference.set(qn("w:id"), note_id)
    run.append(reference)
    paragraph._p.append(run)


def _inject_footnotes(docx_path: Path) -> None:
    footnotes_xml = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:footnotes {nsdecls("w", "m")}>'
        '<w:footnote w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="2">'
        '<w:p><w:r><w:footnoteRef/></w:r><w:r><w:t>脚注内容[1]</w:t></w:r>'
        '<m:oMath><m:r><m:t>x+1</m:t></m:r></m:oMath></w:p>'
        '<w:p><w:r><w:t>脚注第二段</w:t></w:r></w:p>'
        '</w:footnote></w:footnotes>'
    ).encode("utf-8")
    replacement = docx_path.with_suffix(".tmp.docx")
    with ZipFile(docx_path, "r") as source, ZipFile(replacement, "w", ZIP_DEFLATED) as target:
        content_types = etree.fromstring(source.read("[Content_Types].xml"))
        if not content_types.xpath(
            './ct:Override[@PartName="/word/footnotes.xml"]',
            namespaces={"ct": CONTENT_TYPES_NS},
        ):
            override = etree.SubElement(content_types, f"{{{CONTENT_TYPES_NS}}}Override")
            override.set("PartName", "/word/footnotes.xml")
            override.set("ContentType", FOOTNOTES_CONTENT_TYPE)
        relationships = etree.fromstring(source.read("word/_rels/document.xml.rels"))
        relation = etree.SubElement(
            relationships, f"{{{RELATIONSHIPS_NS}}}Relationship"
        )
        relation.set("Id", "rIdScholarFootnotes")
        relation.set("Type", FOOTNOTES_REL_TYPE)
        relation.set("Target", "footnotes.xml")
        for item in source.infolist():
            if item.filename in {
                "[Content_Types].xml", "word/_rels/document.xml.rels", "word/footnotes.xml"
            }:
                continue
            target.writestr(item, source.read(item.filename))
        target.writestr(
            "[Content_Types].xml",
            etree.tostring(content_types, xml_declaration=True, encoding="UTF-8", standalone=True),
        )
        target.writestr(
            "word/_rels/document.xml.rels",
            etree.tostring(relationships, xml_declaration=True, encoding="UTF-8", standalone=True),
        )
        target.writestr("word/footnotes.xml", footnotes_xml)
    shutil.move(replacement, docx_path)


def _write_bib(path: Path) -> None:
    path.write_text(
        "@article{1, author={A}, title={Test}, journal={J}, year={2024}}\n",
        encoding="utf-8",
    )


##### 正常语义链路板块 #####


class WordSemanticPassTests(unittest.TestCase):
    def test_footnote_fields_cross_reference_link_and_text_box_pass(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            docx = root / "semantic.docx"
            bib = root / "cite.bib"
            document = Document()
            document.add_heading("第一章 绪论", level=1)
            target = document.add_paragraph("目标段落")
            _append_bookmark(target, "Target_A")
            reference = document.add_paragraph("参见")
            _append_complex_field(reference, " REF Target_A \\h ", "目标段落")
            reference.add_run("，位于第")
            _append_simple_field(reference, " PAGEREF Target_A \\h ", "1")
            reference.add_run("页。")
            link = document.add_paragraph("官方网站：")
            _append_hyperlink(link, "示例链接", "https://example.com/a?q=1")
            note = document.add_paragraph("带脚注的正文")
            _append_footnote_reference(note, "2")
            box = document.add_paragraph()
            _append_text_box(box, "文本框内容")
            document.save(docx)
            _inject_footnotes(docx)
            _write_bib(bib)

            runner = PipelineRunner()
            extract = runner.extract(str(docx))
            recognize = runner.recognize(extract)
            output = root / "output"
            render = runner.render(
                extract, recognize, str(GRADUATE_TEMPLATE),
                str(output), str(bib),
            )
            report = build_content_fidelity_report(extract, recognize, render)
            tex = (output / render.chapter_files[0]).read_text(encoding="utf-8")
            counts = {
                kind: sum(unit["unit_type"] == kind for unit in extract.content_units)
                for kind in ("footnote", "field", "bookmark", "hyperlink", "text_box")
            }

        self.assertEqual(counts, {
            "footnote": 1, "field": 3, "bookmark": 1,
            "hyperlink": 1, "text_box": 1,
        })
        self.assertTrue(report["all_passed"], report["blocking_or_unimplemented"])
        self.assertEqual(report["metrics"]["footnote_coverage"]["status"], "pass")
        self.assertEqual(report["metrics"]["field_semantic_coverage"]["status"], "pass")
        self.assertEqual(report["metrics"]["cross_reference_target_coverage"]["status"], "pass")
        self.assertEqual(report["metrics"]["text_box_layout_equivalence"]["status"], "degraded")
        self.assertEqual(report["metrics"]["plain_text_integrity"]["status"], "pass")
        self.assertEqual(report["metrics"]["hyperlink_coverage"]["status"], "pass")
        self.assertIn(r"\footnote{脚注内容\cite{1}$x+1$\par 脚注第二段}", tex)
        self.assertIn(r"\hyperref[word-bm-", tex)
        self.assertIn(r"\pageref{word-bm-", tex)
        self.assertIn(r"\href{https://example.com/a?q=1}{示例链接}", tex)
        self.assertIn(r"\fbox{\parbox{0.92\linewidth}{文本框内容$y+1$", tex)
        self.assertIn(r"\href{https://example.org}{框内链接}", tex)


##### 明确失败板块 #####


class WordSemanticFailureTests(unittest.TestCase):
    def test_missing_bookmark_and_malformed_field_fail(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            docx = root / "field-fail.docx"
            bib = root / "cite.bib"
            document = Document()
            document.add_heading("第一章 绪论", level=1)
            missing = document.add_paragraph("缺失目标：")
            _append_complex_field(missing, " REF Missing_Target \\h ", "目标")
            malformed = document.add_paragraph("畸形域：")
            _append_malformed_field(malformed, " REF Missing_Target ", "目标")
            document.save(docx)
            _write_bib(bib)
            runner = PipelineRunner()
            extract = runner.extract(str(docx))
            recognize = runner.recognize(extract)
            output = root / "output"
            render = runner.render(
                extract, recognize, str(GRADUATE_TEMPLATE),
                str(output), str(bib),
            )
            report = build_content_fidelity_report(extract, recognize, render)
            issue_codes = {
                item["code"] for item in extract.extraction_report["issues"]
            }
        self.assertFalse(report["all_passed"])
        self.assertEqual(report["metrics"]["field_semantic_coverage"]["status"], "fail")
        self.assertEqual(report["metrics"]["cross_reference_target_coverage"]["status"], "fail")
        self.assertIn("FIELD_STRUCTURE_MALFORMED", issue_codes)

    def test_missing_footnote_unsupported_field_unsafe_link_and_complex_box_fail(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            docx = root / "semantic-fail.docx"
            bib = root / "cite.bib"
            document = Document()
            document.add_heading("第一章 绪论", level=1)
            note = document.add_paragraph("缺失脚注")
            _append_footnote_reference(note, "99")
            field = document.add_paragraph("域：")
            _append_simple_field(field, " ADDIN Zotero.Item CSL_CITATION ", "静态引文")
            link = document.add_paragraph("本地链接：")
            _append_hyperlink(link, "危险链接", "file:///C:/secret.txt")
            box = document.add_paragraph()
            _append_text_box(box, "复杂文本框", nested_table=True)
            document.save(docx)
            _write_bib(bib)

            runner = PipelineRunner()
            extract = runner.extract(str(docx))
            recognize = runner.recognize(extract)
            output = root / "output"
            render = runner.render(
                extract, recognize, str(GRADUATE_TEMPLATE),
                str(output), str(bib),
            )
            report = build_content_fidelity_report(extract, recognize, render)
            issue_codes = {
                item["code"] for item in extract.extraction_report["issues"]
            }

        self.assertFalse(report["all_passed"])
        self.assertEqual(report["metrics"]["footnote_coverage"]["status"], "fail")
        self.assertEqual(report["metrics"]["field_semantic_coverage"]["status"], "fail")
        self.assertEqual(report["metrics"]["hyperlink_coverage"]["status"], "fail")
        self.assertEqual(report["metrics"]["text_box_content_coverage"]["status"], "fail")
        self.assertTrue({
            "FOOTNOTE_TARGET_MISSING",
            "FIELD_SEMANTIC_UNSUPPORTED",
            "HYPERLINK_TARGET_UNSAFE_OR_MISSING",
            "TEXT_BOX_COMPLEX_CONTENT_UNSUPPORTED",
        }.issubset(issue_codes))


if __name__ == "__main__":
    unittest.main(verbosity=2)
