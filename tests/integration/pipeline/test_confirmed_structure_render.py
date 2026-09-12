# -*- coding: utf-8 -*-
"""确认标题层级必须真正改变 LaTeX，而不是只修改后端 JSON。"""

from __future__ import annotations

import copy
import hashlib
from pathlib import Path

import pytest
from docx import Document

from pipeline.confirmed_structure import ConfirmedStructureError, apply_confirmed_structure
from pipeline.confirmed_structure import build_structure_snapshot, finalize_structure_snapshot
from pipeline_api import PipelineRunner
from tests.support.paths import GRADUATE_TEMPLATE


##### 虚拟确认快照板块 #####


def _prepare(tmp_path: Path):
    docx = tmp_path / "paper.docx"
    document = Document()
    document.add_heading("第一章 绪论", level=1)
    document.add_paragraph("章节正文。")
    document.add_heading("1.1 应降为正文的标题", level=2)
    document.add_paragraph("后续正文。")
    document.save(docx)
    runner = PipelineRunner()
    extracted = runner.extract(str(docx))
    recognized = runner.recognize(extracted)
    digest = hashlib.sha256(docx.read_bytes()).hexdigest()
    snapshot = build_structure_snapshot(extracted, recognized, digest)
    for heading in snapshot["headings"]:
        if heading["level"] == "section":
            heading.update(level="body", overridden=True, review_status="resolved")
    snapshot = finalize_structure_snapshot(snapshot)
    snapshot["revision"] = 1
    return runner, docx, extracted, recognized, snapshot, digest


##### 应用与渲染板块 #####


def test_confirmed_body_level_changes_rendered_latex(tmp_path: Path) -> None:
    runner, _, extracted, recognized, snapshot, digest = _prepare(tmp_path)
    confirmed = apply_confirmed_structure(
        extracted, recognized, snapshot,
        source_sha256=digest, structure_revision=1,
    )
    rendered = runner.render(
        extracted, confirmed, str(GRADUATE_TEMPLATE),
        str(tmp_path / "render"), None,
    )
    section = (Path(rendered.output_dir) / "contents" / "section_01.tex").read_text(encoding="utf-8")
    assert r"\section{应降为正文的标题}" not in section
    assert "1.1 应降为正文的标题" in section
    assert confirmed.internal["structure_revision"] == 1


def test_confirmed_structure_rejects_wrong_source_hash(tmp_path: Path) -> None:
    _, _, extracted, recognized, snapshot, _ = _prepare(tmp_path)
    with pytest.raises(ConfirmedStructureError, match="哈希不一致"):
        apply_confirmed_structure(
            extracted, recognized, snapshot,
            source_sha256="wrong", structure_revision=1,
        )


@pytest.mark.parametrize("decision", ["body_citation", "non_citation"])
def test_confirmed_citations_change_generation_without_editing_content(tmp_path, decision):
    from content_extraction.citations import apply_citation_overrides
    from tests.unit.content_extraction.test_citations import citation_document
    from docx.oxml.ns import qn

    document = citation_document()
    for paragraph in list(document.paragraphs):
        if paragraph._p.find(".//" + qn("w:footnoteReference")) is not None:
            document.element.body.remove(paragraph._p)
    path = tmp_path / "citations.docx"
    document.save(path)
    runner = PipelineRunner()
    extracted = runner.extract(str(path))
    recognized = runner.recognize(extracted)
    original = [u.copy() for u in extracted.content_units]
    ambiguous = next(item for item in recognized.review["citation_review"]["decisions"] if item["decision"] == "needs_review")
    digest = extracted.extraction_report["metadata"]["source_sha256"]
    snapshot = build_structure_snapshot(extracted, recognized, digest)
    snapshot["citation_review"] = apply_citation_overrides(recognized.review["citation_review"], [
        {"unit_id": ambiguous["unit_id"], "decision": decision},
    ])
    snapshot = finalize_structure_snapshot(snapshot)
    snapshot["revision"] = 1
    confirmed = apply_confirmed_structure(extracted, recognized, snapshot, source_sha256=digest, structure_revision=1)
    rendered = runner.render(extracted, confirmed, str(GRADUATE_TEMPLATE), str(tmp_path / "render"), None, reference_source="word")
    content = (Path(rendered.output_dir) / rendered.chapter_files[0]).read_text(encoding="utf-8")
    assert r"已有研究\cite{word-ref-0001,word-ref-0002,word-ref-0003}。" in content
    expected = r"符号\cite{word-ref-0001}位于此处" if decision == "body_citation" else "符号[1]位于此处"
    assert expected in content
    assert extracted.content_units == original
    assert next(item for item in recognized.review["citation_review"]["decisions"] if item["unit_id"] == ambiguous["unit_id"])["decision"] == "needs_review"
