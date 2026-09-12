# -*- coding: utf-8 -*-
"""G4-A 图片、复杂表格和 Word 原生公式的虚拟案例。"""

from __future__ import annotations

import base64
import hashlib
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement

from tests.support.paths import STAGE1_DIR


##### 测试环境板块 #####


from adapters.ouc import OUCTemplateAdapter
from pipeline.chapter_renderer import render_chapter_files
from pipeline.content_fidelity_gate import build_content_fidelity_report
from pipeline.object_renderer import materialize_word_assets
from pipeline.results import RenderResult
from pipeline_api import PipelineRunner


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUB"
    "AScY42YAAAAASUVORK5CYII="
)


def _append_omml(paragraph, text: str = "x+1") -> None:
    formula = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    value = OxmlElement("m:t")
    value.text = text
    run.append(value)
    formula.append(run)
    paragraph._p.append(formula)


def _build_virtual_thesis(root: Path) -> Path:
    image_path = root / "pixel.png"
    image_path.write_bytes(PNG_1X1)
    path = root / "g4-virtual.docx"
    document = Document()
    document.add_heading("第一章 绪论", level=1)
    formula_paragraph = document.add_paragraph("由公式")
    _append_omml(formula_paragraph)
    formula_paragraph.add_run("可得结果。")
    document.add_picture(str(image_path))
    document.add_paragraph("图1-1 内嵌图片")
    document.add_paragraph("表1-1 合并单元格")
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并表头"
    table.cell(0, 2).text = "指标"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "分组A"
    table.cell(1, 1).text = "甲"
    table.cell(1, 2).text = "1"
    _append_omml(table.cell(1, 2).paragraphs[0], "a+b")
    table.cell(2, 1).text = "乙"
    table.cell(2, 2).text = "2"
    table.cell(2, 2).paragraphs[0].add_run().add_picture(str(image_path))
    document.add_paragraph("正文结束。")
    document.save(path)
    return path


##### 端到端对象测试板块 #####


class WordObjectPipelineTests(unittest.TestCase):
    def test_embedded_image_merge_and_omml_pass_fidelity(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            docx_path = _build_virtual_thesis(root)
            runner = PipelineRunner()
            extract = runner.extract(str(docx_path))
            recognize = runner.recognize(extract)

            image = next(unit for unit in extract.content_units if unit["unit_type"] == "image")
            self.assertEqual(
                image["payload"]["binary_sha256"], hashlib.sha256(PNG_1X1).hexdigest()
            )
            self.assertEqual(extract.media_assets[image["unit_id"]], PNG_1X1)

            paragraph = next(
                unit for unit in extract.content_units
                if unit["unit_type"] == "paragraph" and unit["text"].startswith("由公式")
            )
            self.assertEqual(
                [token["kind"] for token in paragraph["payload"]["inline_tokens"]],
                ["text", "formula", "text"],
            )
            table = next(
                unit for unit in extract.content_units
                if unit["unit_type"] == "table"
            )
            cells = [cell for row in table["payload"]["physical_rows"] for cell in row["cells"]]
            self.assertTrue(any(cell["grid_span"] == 2 for cell in cells))
            self.assertTrue(any(cell["vertical_merge"] == "restart" for cell in cells))
            self.assertTrue(any(cell["vertical_merge"] == "continue" for cell in cells))

            contents = root / "output" / "contents"
            contents.mkdir(parents=True)
            manifest = materialize_word_assets(
                extract.content_units, extract.media_assets, root / "output"
            )
            stats = {"figures": 0, "tables": 0, "equations": 0, "citations": 0}
            chapter_files, trace = render_chapter_files(
                template_adapter=OUCTemplateAdapter(), content_directory="contents",
                chapters=recognize.review["chapters"],
                object_bindings=recognize.review["object_bindings"],
                paragraphs=extract.paragraphs,
                content_units=extract.content_units,
                cite_map={},
                contents_dir=contents,
                stats=stats,
                source_schema_version="1.3.0",
                asset_manifest=manifest,
                confirmed_headings={item["unit_id"]: item for item in recognize.review["heading_candidates"]},
            )
            render = RenderResult(
                template_adapter=OUCTemplateAdapter(),
                output_dir=str(root / "output"),
                stats=stats,
                chapter_files=chapter_files,
                render_trace=trace,
            )
            report = build_content_fidelity_report(extract, recognize, render)
            tex = (root / "output" / chapter_files[0]).read_text(encoding="utf-8")

        self.assertTrue(report["all_passed"], report["blocking_or_unimplemented"])
        self.assertEqual(report["metrics"]["image_object_coverage"]["status"], "pass")
        self.assertEqual(report["metrics"]["table_cell_coverage"]["status"], "pass")
        self.assertEqual(report["metrics"]["omml_formula_coverage"]["status"], "pass")
        self.assertIn("由公式$x+1$可得结果", tex)
        self.assertIn(r"\multicolumn{2}{c}{合并表头}", tex)
        self.assertIn(r"\multirow{2}{*}{分组A}", tex)

    def test_nested_table_is_recorded_as_unsupported(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "nested.docx"
            document = Document()
            document.add_heading("第一章 绪论", level=1)
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "嵌套"
            document.save(path)
            extract = PipelineRunner().extract(str(path))
        table_unit = next(unit for unit in extract.content_units if unit["unit_type"] == "table")
        self.assertEqual(table_unit["status"], "degraded")
        self.assertIn("nested_table", table_unit["properties"]["unsupported_features"])
        issue_codes = [item["code"] for item in extract.extraction_report["issues"]]
        self.assertIn("TABLE_PRESENTATION_UNSUPPORTED", issue_codes)


if __name__ == "__main__":
    unittest.main(verbosity=2)
