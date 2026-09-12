# -*- coding: utf-8 -*-
"""首批合成 Word 输入本身的可重复性和人工规格检查。"""

from __future__ import annotations

import hashlib
import json

import pytest
from docx import Document

from tests.support.build_word_cases import REPOSITORY_ROOT, _semantic_sha256


##### 案例加载板块 #####


BASELINE_DIR = REPOSITORY_ROOT / "tests" / "baselines" / "synthetic"


def _expectation(case_name: str) -> dict:
    return json.loads(
        (BASELINE_DIR / f"{case_name}.expected.json").read_text(encoding="utf-8")
    )


##### 通用完整性板块 #####


@pytest.mark.parametrize(
    "case_name",
    [
        "W01_minimal_body", "W03_ambiguous_structure", "W05_references_complete",
        "W07_formulas_mixed_valid", "W09_formulas_malformed",
    ],
)
def test_synthetic_word_matches_recorded_hashes(case_name: str) -> None:
    expectation = _expectation(case_name)
    path = REPOSITORY_ROOT / expectation["artifact"]["relative_path"]
    assert path.is_file()
    assert hashlib.sha256(path.read_bytes()).hexdigest() == expectation["artifact"][
        "package_sha256"
    ]
    assert _semantic_sha256(Document(path)) == expectation["artifact"][
        "semantic_sha256"
    ]
    assert expectation["source_kind"] == "fully_synthetic"


##### 案例语义板块 #####


def test_w01_really_contains_only_minimal_text() -> None:
    document = Document(
        REPOSITORY_ROOT / "tests" / "fixtures" / "word" / "W01_minimal_body.docx"
    )
    assert len(document.paragraphs) == 6
    assert not document.tables
    assert [item.style.name for item in document.paragraphs] == [
        "Title", "Heading 1", "Normal", "Normal", "Heading 2", "Normal",
    ]


def test_w03_first_table_is_obviously_content_data() -> None:
    document = Document(
        REPOSITORY_ROOT / "tests" / "fixtures" / "word" / "W03_ambiguous_structure.docx"
    )
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == "监测时刻"
    assert document.tables[0].cell(0, 1).text == "温度"
    noisy = {paragraph.text: paragraph.style.name for paragraph in document.paragraphs}
    assert noisy["1 绪论"] == "Normal"
    assert noisy["研究方法"] == "Normal"
    assert noisy["实验结果"] == "Normal"


def test_w05_keeps_explicit_reference_numbers_and_context_examples() -> None:
    document = Document(
        REPOSITORY_ROOT / "tests" / "fixtures" / "word" / "W05_references_complete.docx"
    )
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "已有研究给出了相关结果[1]。" in texts
    assert "另一项分析对该方法进行了扩展[2-3]。" in texts
    assert "归一化区间为[0,1]，这里不是参考文献引用。" in texts
    assert "数组索引示例为[1,2]，这里也不是参考文献引用。" in texts
    references = [text for text in texts if text.startswith(("[1]", "[2]", "[3]"))]
    assert len(references) == 3


def test_w07_contains_two_outer_omml_formulas_and_two_latex_formulas() -> None:
    path = REPOSITORY_ROOT / "tests" / "fixtures" / "word" / "W07_formulas_mixed_valid.docx"
    document = Document(path)
    outer = document.element.xpath(
        ".//*[local-name()='oMathPara'] | "
        ".//*[local-name()='oMath' and not(parent::*[local-name()='oMathPara'])]"
    )
    assert len(outer) == 2
    assert len(document.element.xpath(".//*[local-name()='f']")) == 1
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "正文中的 LaTeX 行内公式 $a^2+b^2=c^2$ 应保持行内。" in texts
    assert r"$$\int_0^1 x^2\,dx=\frac{1}{3}$$" in texts
    assert "归一化区间为[0,1]，不能识别为引用或公式。" in texts
    assert r"预算写作 \$100 时，转义美元符号不能开启公式。" in texts


def test_w09_preserves_three_malformed_formula_candidates_as_raw_text() -> None:
    document = Document(
        REPOSITORY_ROOT / "tests" / "fixtures" / "word" / "W09_formulas_malformed.docx"
    )
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "未闭合的行内公式为 $x+y。" in texts
    assert "单双美元混用的公式为 $$x+y$。" in texts
    assert r"花括号未闭合的公式为 $$\frac{a}{b$$。" in texts
    assert texts[-1] == "错误公式之后的正文必须继续保留。"
