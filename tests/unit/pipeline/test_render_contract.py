# -*- coding: utf-8 -*-
"""渲染与质量门禁整改的独立故障注入。"""

import re
from copy import deepcopy
from pathlib import Path
import pytest
from docx import Document

from content_extraction.docx_extractor import extract_docx_ordered
from models.render_trace import fragment_hash
from pipeline.content_fidelity_gate import build_content_fidelity_report
from pipeline.text_processing import escape_cell, escape_text
from pipeline.object_renderer import render_table, render_figure
from pipeline.text_processing import hyperlink_argument
from adapters.ouc_bachelor import OUCBachelorTemplateAdapter
from adapters.ouc import OUCTemplateAdapter, CHAPTERS_BEGIN, CHAPTERS_END
from pipeline.structural_gate import structural_fidelity_gate, FatalStructuralMismatch
from pipeline_api import PipelineRunner
from pipeline.chapter_renderer import render_chapter_files
from pipeline.results import RenderResult
from tests.unit.pipeline.test_content_fidelity_gate import _virtual_case
from pipeline.table_pdf_check import table_row_signatures, table_rows
from pipeline.object_renderer import _marked_cell


##### 普通文字契约板块 #####


def test_all_plain_special_characters_are_escaped():
    assert escape_text('#$%&~_^\\{}') == (
        r'\#\$\%\&\textasciitilde{}\_\textasciicircum{}\textbackslash{}\{\}'
    )


def test_complete_url_is_recognized_before_surrounding_escaping():
    source = '地址 https://example.org/a_b?q=x&other=50%25#part_1。费用 $5。'
    assert escape_text(source) == (
        r'地址 \url{https://example.org/a_b?q=x&other=50%25#part_1}。费用 \$5。'
    )


def test_explicit_hyperlink_preserves_query_percent_encoding_and_fragment():
    assert hyperlink_argument('https://example.org/a?q=20%25&x=1#part') == r'https://example.org/a?q=20\%25\&x=1\#part'


def test_whitespace_and_unicode_are_shared_with_cells():
    assert escape_text(' \t ') == ' '
    assert escape_text('') == ''
    assert escape_cell('α ≥ Ⅱ_2') == escape_text('α ≥ Ⅱ_2')


def test_cell_formulas_and_citations_are_extracted_as_semantic_tokens(tmp_path):
    path = tmp_path / 'nested.docx'
    document = Document()
    document.add_paragraph('1 绪论')
    document.add_table(rows=1, cols=1).cell(0, 0).text = '结果 $x_1$，参考[1]。'
    document.save(path)
    bundle = extract_docx_ordered(str(path))
    table = next(unit for unit in bundle.units_as_dicts() if unit['unit_type'] == 'table')
    tokens = table['payload']['physical_rows'][0]['cells'][0]['tokens']
    assert [token['kind'] for token in tokens] == ['text', 'formula', 'text', 'citation', 'text']


def test_manual_break_before_formula_preserves_token_offsets(tmp_path):
    document = Document()
    paragraph = document.add_paragraph('第一行')
    paragraph.add_run().add_break()
    paragraph.add_run('公式 $x_1$。')
    path = tmp_path / 'line-break.docx'
    document.save(path)
    units = extract_docx_ordered(str(path)).units_as_dicts()
    tokens = next(unit for unit in units if unit['unit_type'] == 'paragraph')['payload']['inline_tokens']
    assert [token['kind'] for token in tokens] == ['text', 'line_break', 'text', 'formula', 'text']


def test_multiline_cell_markers_preserve_logical_rows_for_pdf_checks():
    rows = '\n'.join(_marked_cell(f'u-000020:r{i}:c0', value) + r' \\'
                     for i, value in enumerate(('firstvalue', 'secondvalue')))
    source = '% SCHOLAR_UNIT_BEGIN u-000020 table_caption\n' + r'\begin{longtable}{c}' + '\n'
    source += r'\endlastfoot' + '\n' + rows + '\n' + r'\end{longtable}'
    source += '\n% SCHOLAR_UNIT_END u-000020'
    assert len(table_rows(rows)) == 2
    assert table_row_signatures(source, 'u-000020') == ['firstvalue', 'secondvalue']


def test_table_rejects_nonempty_vertical_continuation_without_mutation():
    table = {'unit_id': 'u-000020', 'properties': {'columns': 1, 'unsupported_features': []},
             'payload': {'physical_rows': [{'cells': [
                 {'cell_id': 'u-000020:r0:c0', 'grid_column': 0, 'vertical_merge': 'continue',
                  'diagonal_border': 'none', 'text': '不得丢弃',
                  'tokens': [{'kind': 'text', 'text': '不得丢弃'}]},
             ]}]}}
    original = deepcopy(table)
    _, details, rendered = render_table('表1', '题注', table, {}, template_adapter=OUCBachelorTemplateAdapter())
    assert not rendered
    assert details['unsupported_features']
    assert table == original


def test_bachelor_caption_uses_its_supported_command():
    adapter = OUCBachelorTemplateAdapter()
    assert adapter.render_caption('figure', '图片', '') == r'\caption{图片}'


def test_duplicate_formula_occurrence_in_ledger_blocks(tmp_path):
    extracted, recognized, rendered, _ = _virtual_case(tmp_path)
    record = next(item for item in rendered.render_trace['records'] if item['marker_unit_id'] == 'u-000001')
    record['details']['rendered_formula_ids'].append('u-000002')
    assert not build_content_fidelity_report(extracted, recognized, rendered)['all_passed']


def test_duplicate_image_occurrence_is_not_hidden_by_set(tmp_path):
    extracted, recognized, rendered, _ = _virtual_case(tmp_path, with_image=True)
    record = next(item for item in rendered.render_trace['records'] if item['marker_unit_id'] == 'u-000005')
    record['details']['verified_image_unit_ids'] = ['u-000007', 'u-000007']
    assert not build_content_fidelity_report(extracted, recognized, rendered)['all_passed']


def test_figure_layout_cannot_duplicate_a_confirmed_image():
    binding = {'object_unit_ids': ['u-000007'], 'layout_rows': [['u-000007', 'u-000007']]}
    assets = {'u-000007': {'status': 'verified', 'relative_path': 'figures/image.png'}}
    _, _, rendered = render_figure('图1', '图片', binding, assets, template_adapter=OUCTemplateAdapter())
    assert not rendered


def test_cell_text_deletion_with_synced_ledger_blocks(tmp_path):
    document = Document()
    document.add_heading('第一章 绪论', level=1)
    document.add_table(rows=1, cols=1).cell(0, 0).text = '单元格原文'
    path = tmp_path / 'table.docx'
    document.save(path)
    runner = PipelineRunner()
    extracted = runner.extract(str(path))
    recognized = runner.recognize(extracted)
    contents = tmp_path / 'contents'
    contents.mkdir()
    stats = {'figures': 0, 'tables': 0, 'equations': 0, 'citations': 0}
    files, trace = render_chapter_files(chapters=recognized.review['chapters'],
        object_bindings=[], paragraphs=extracted.paragraphs, content_units=extracted.content_units,
        cite_map={}, contents_dir=contents, stats=stats, source_schema_version='1.7.0',
        template_adapter=OUCTemplateAdapter(), content_directory='contents',
        confirmed_headings={item['unit_id']: item for item in recognized.review['heading_candidates']})
    rendered = RenderResult(str(tmp_path), stats, files, trace, template_adapter=OUCTemplateAdapter())
    assert build_content_fidelity_report(extracted, recognized, rendered)['all_passed']
    latex_path = tmp_path / files[0]
    content = latex_path.read_text(encoding='utf-8').replace('单元格原文', '单元格')
    latex_path.write_text(content, encoding='utf-8')
    record = next(item for item in trace['records'] if item['role'] == 'content_table')
    fragment = re.search(r'% SCHOLAR_UNIT_BEGIN ' + record['marker_unit_id'] + r' content_table\n(.*?)\n% SCHOLAR_UNIT_END', content, re.S).group(1)
    record['output_hash'] = fragment_hash(fragment)
    assert not build_content_fidelity_report(extracted, recognized, rendered)['all_passed']


##### 同步篡改账本板块 #####


@pytest.mark.parametrize('fault', ['append', 'remove'])
def test_english_heading_requires_one_confirmed_source_and_translation_slot(tmp_path, fault):
    extracted, recognized, rendered, path = _virtual_case(tmp_path)
    text = path.read_text(encoding='utf-8')
    if fault == 'append':
        text += '\n' + r'\enchapter{Unsourced heading}'
    else:
        text = text.replace(r'\enchapter{绪论}', '')
    path.write_text(text, encoding='utf-8')
    assert not build_content_fidelity_report(extracted, recognized, rendered)['all_passed']


@pytest.mark.parametrize('fault', ['role', 'source', 'target', 'hash'])
def test_ledger_role_source_and_target_are_verified_independently(tmp_path, fault):
    extracted, recognized, rendered, _ = _virtual_case(tmp_path)
    record = next(item for item in rendered.render_trace['records'] if item['marker_unit_id'] == 'u-000005')
    if fault == 'role':
        record['role'] = 'chapter_heading'
    elif fault == 'source':
        record['source_unit_ids'].append('u-999999')
    elif fault == 'target':
        record['target_file'] = 'contents/section_99.tex'
    else:
        record['output_hash'] = '0' * 64
    assert not build_content_fidelity_report(extracted, recognized, rendered)['all_passed']


@pytest.mark.parametrize('fault', ['swap_levels', 'comment_heading'])
def test_individual_heading_contract_cannot_be_replaced_by_counts(tmp_path, fault):
    extracted, recognized, rendered, path = _virtual_case(tmp_path)
    recognized.review['heading_candidates'] = [
        {'unit_id': 'u-000000', 'level': 'chapter', 'title': '绪论'},
        {'unit_id': 'u-000004', 'level': 'section', 'title': '研究方法'},
    ]
    main = tmp_path / 'main.tex'
    main.write_text(CHAPTERS_BEGIN + '\n' + CHAPTERS_END, encoding='utf-8')
    adapter = OUCTemplateAdapter()
    adapter.update_chapter_includes(main, rendered.chapter_files)
    content = path.read_text(encoding='utf-8')
    if fault == 'swap_levels':
        content = content.replace(r'\chapter{绪论}', r'\section{绪论}').replace(r'\section{研究方法}', r'\chapter{研究方法}')
    else:
        content = content.replace(r'\chapter{绪论}', r'%\chapter{绪论}')
    path.write_text(content, encoding='utf-8')
    with pytest.raises(FatalStructuralMismatch):
        structural_fidelity_gate(extracted.word_structure, tmp_path / 'contents',
            confirmed_structure=recognized.review, main_tex_path=main,
            chapter_files=rendered.chapter_files, template_adapter=adapter)


def test_deleted_body_text_with_synchronized_ledger_still_blocks(tmp_path: Path):
    extract, recognized, rendered, path = _virtual_case(tmp_path)
    content = path.read_text(encoding='utf-8')
    content = content.replace('重复内容。', '内容。', 1)
    path.write_text(content, encoding='utf-8')
    fragment = re.search(
        r'% SCHOLAR_UNIT_BEGIN u-000005 body_paragraph\n(.*?)\n% SCHOLAR_UNIT_END',
        content, re.S,
    ).group(1)
    record = next(item for item in rendered.render_trace['records']
                  if item['marker_unit_id'] == 'u-000005')
    record['output_hash'] = fragment_hash(fragment)
    report = build_content_fidelity_report(extract, recognized, rendered)
    assert not report['all_passed']


def test_inserted_unsourced_text_inside_marker_with_synced_ledger_blocks(tmp_path: Path):
    extract, recognized, rendered, path = _virtual_case(tmp_path)
    content = path.read_text(encoding='utf-8')
    content = content.replace('重复内容。', '重复内容。无来源文字。', 1)
    path.write_text(content, encoding='utf-8')
    fragment = re.search(
        r'% SCHOLAR_UNIT_BEGIN u-000005 body_paragraph\n(.*?)\n% SCHOLAR_UNIT_END',
        content, re.S,
    ).group(1)
    record = next(item for item in rendered.render_trace['records']
                  if item['marker_unit_id'] == 'u-000005')
    record['output_hash'] = fragment_hash(fragment)
    assert not build_content_fidelity_report(extract, recognized, rendered)['all_passed']
