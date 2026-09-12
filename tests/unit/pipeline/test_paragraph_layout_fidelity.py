# -*- coding: utf-8 -*-
"""G9-C1 段落布局保真和公式解释段落专项测试。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from adapters.ouc import OUCTemplateAdapter
from pipeline.chapter_renderer import render_chapter_files
from content_extraction.docx_extractor import extract_docx_ordered


##### 虚拟文档板块 #####


def _build_layout_case(path: Path) -> None:
    document = Document()
    document.add_heading("1 绪论", level=1)
    document.add_paragraph("$x+y$")
    explanation = document.add_paragraph("其中，x 表示输入。")
    explanation.paragraph_format.first_line_indent = Pt(0)
    ordinary = document.add_paragraph("其中，普通正文不能仅按关键词判断。")
    ordinary.paragraph_format.first_line_indent = Pt(0)
    document.add_paragraph("$a+b$")
    indented = document.add_paragraph("其中，该段在 Word 中明确缩进。")
    indented.paragraph_format.first_line_indent = Pt(24)
    document.save(path)


##### 抽取与渲染板块 #####


def test_equation_explanation_requires_display_adjacency_and_zero_indent(
    tmp_path: Path,
) -> None:
    source = tmp_path / "layout.docx"
    _build_layout_case(source)
    bundle = extract_docx_ordered(str(source))
    paragraphs = [unit for unit in bundle.units if unit.unit_type == "paragraph"]

    assert paragraphs[2].properties["layout"]["first_line_indent"]["state"] == "zero"
    assert paragraphs[2].properties["semantic_role"] == "equation_explanation"
    assert paragraphs[3].properties.get("semantic_role") is None
    assert paragraphs[5].properties["layout"]["first_line_indent"]["state"] == "first_line"
    assert paragraphs[5].properties.get("semantic_role") is None


def test_adapter_renders_only_confirmed_equation_explanation_without_indent(
    tmp_path: Path,
) -> None:
    source = tmp_path / "layout.docx"
    _build_layout_case(source)
    bundle = extract_docx_ordered(str(source))
    contents = tmp_path / "contents"
    contents.mkdir()
    stats = {"figures": 0, "tables": 0, "equations": 0, "citations": 0}

    chapter_files, trace = render_chapter_files(
        template_adapter=OUCTemplateAdapter(), content_directory="contents",
        chapters=[{"num": 1, "title": "绪论", "range": (0, len(bundle.paragraphs))}],
        object_bindings=[], paragraphs=bundle.paragraphs,
        content_units=bundle.units_as_dicts(), cite_map={},
        contents_dir=contents, stats=stats,
        source_schema_version=bundle.schema_version,
        confirmed_headings={bundle.paragraphs[0]["unit_id"]: {"level": "chapter", "title": "绪论"}},
    )
    content = (tmp_path / chapter_files[0]).read_text(encoding="utf-8")

    assert r"\noindent 其中，x 表示输入。" in content
    assert r"\noindent 其中，普通正文不能仅按关键词判断。" not in content
    assert len(trace["headings"]) == 1
    assert r"\enchapter" not in Path(
        "stage1/pipeline/chapter_renderer.py"
    ).read_text(encoding="utf-8")
