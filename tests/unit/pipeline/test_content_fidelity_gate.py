# -*- coding: utf-8 -*-
"""G3-A 内容单元、渲染追踪和正文保真故障注入测试。"""

from __future__ import annotations

import re
import tempfile
import unittest
from pathlib import Path

from docx import Document

from tests.support.paths import STAGE1_DIR


##### 测试环境板块 #####


from content_extraction.content_units import ContentUnit
from adapters.ouc import OUCTemplateAdapter
from pipeline.chapter_renderer import render_chapter_files
from pipeline.content_fidelity_gate import build_content_fidelity_report
from content_extraction.docx_extractor import extract_docx_ordered
from pipeline.results import ExtractResult, RecognizeResult, RenderResult


def _unit(unit_id, unit_type, order, text="", **kwargs):
    return ContentUnit(
        unit_id=unit_id,
        unit_type=unit_type,
        order=order,
        text=text,
        **kwargs,
    ).to_dict()


def _virtual_case(root: Path, *, with_image: bool = False, unresolved: bool = False):
    paragraphs = [
        {"index": 0, "block_index": 0, "unit_id": "u-000000", "style": "Heading 1", "text": "1 绪论", "outline_level": 0, "numbering_level": 0},
        {"index": 1, "block_index": 1, "unit_id": "u-000001", "style": "Normal", "text": "结果为 $x+y$，参考[1]。", "outline_level": None, "numbering_level": None},
        {"index": 2, "block_index": 2, "unit_id": "u-000004", "style": "Heading 2", "text": "1.1 研究方法", "outline_level": 1, "numbering_level": 1},
        {"index": 3, "block_index": 3, "unit_id": "u-000005", "style": "Normal", "text": "重复内容。", "outline_level": None, "numbering_level": None},
        {"index": 4, "block_index": 4, "unit_id": "u-000006", "style": "Normal", "text": "重复内容。", "outline_level": None, "numbering_level": None},
    ]
    units = [
        _unit("u-000000", "paragraph", 0, "1 绪论", source={"block_index": 0, "paragraph_index": 0}),
        _unit(
            "u-000001", "paragraph", 1, paragraphs[1]["text"],
            source={"block_index": 1, "paragraph_index": 1},
            payload={"inline_tokens": [
                {"kind": "text", "text": "结果为 "},
                {"kind": "formula", "unit_id": "u-000002"},
                {"kind": "text", "text": "，参考"},
                {"kind": "citation", "unit_id": "u-000003"},
                {"kind": "text", "text": "。"},
            ]},
        ),
        _unit("u-000002", "formula", 2, "x+y", source={"block_index": 1, "start_offset": 4, "end_offset": 9}, properties={"display": False}, relations={"parent_unit_id": "u-000001"}, payload={"source_syntax": "latex", "delimiter": "$"}),
        _unit("u-000003", "citation", 3, "[1]", source={"block_index": 1, "start_offset": 12, "end_offset": 15}, relations={"parent_unit_id": "u-000001"}, payload={"numbers": ["1"]}),
        _unit("u-000004", "paragraph", 4, "1.1 研究方法", source={"block_index": 2, "paragraph_index": 2}),
        _unit("u-000005", "paragraph", 5, "重复内容。", source={"block_index": 3, "paragraph_index": 3}),
        _unit("u-000006", "paragraph", 6, "重复内容。", source={"block_index": 4, "paragraph_index": 4}),
    ]
    if with_image:
        units.append(
            _unit(
                "u-000007",
                "image",
                7,
                source={"block_index": 3, "inline_index": 0},
                relations={"parent_unit_id": "u-000005"},
                payload={"target_part": "media/image1.png"},
            )
        )
    extract = ExtractResult(
        paragraphs=paragraphs,
        tables=[],
        content_units=units,
        extraction_report={"schema_version": "1.1.0", "metadata": {"source_sha256": "0" * 64}},
    )
    recognize = RecognizeResult(
        review={"chapters": [{"num": 1, "title": "绪论", "range": (0, 5)}], "object_bindings": [],
                "citation_review": {"decisions": [{"unit_id": "u-000003", "decision": "body_citation"}]}},
        internal={},
    )
    contents = root / "contents"
    contents.mkdir()
    stats = {"figures": 0, "tables": 0, "equations": 0, "citations": 0}
    cite_map = {} if unresolved else {"1": "1+example"}
    chapter_files, trace = render_chapter_files(
        template_adapter=OUCTemplateAdapter(), content_directory="contents",
        chapters=recognize.review["chapters"],
        object_bindings=[],
        paragraphs=paragraphs,
        content_units=units,
        cite_map=cite_map,
        contents_dir=contents,
        stats=stats,
        source_schema_version="1.7.0",
        confirmed_headings={"u-000000": {"level": "chapter", "title": "绪论"}, "u-000004": {"level": "section", "title": "研究方法"}},
        citation_decisions={"u-000003": {"decision": "body_citation"}},
    )
    trace["bibliography"] = {"render_number_to_key": cite_map,
        "body_citation_unit_ids": ["u-000003"], "all_citations_resolved": not unresolved}
    render = RenderResult(
        template_adapter=OUCTemplateAdapter(),
        output_dir=str(root),
        stats=stats,
        chapter_files=chapter_files,
        render_trace=trace,
    )
    recognize.review["heading_candidates"] = [
        {"unit_id": "u-000000", "level": "chapter", "title": "绪论"},
        {"unit_id": "u-000004", "level": "section", "title": "研究方法"},
    ]
    return extract, recognize, render, root / chapter_files[0]


##### 抽取子单元板块 #####


class TextFragmentExtractionTests(unittest.TestCase):
    def test_latex_formula_citation_hash_and_parent_relation_are_stable(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "fragments.docx"
            document = Document()
            document.add_paragraph("1 绪论")
            document.add_paragraph("公式 $x+y$，参见[1, 2]。")
            document.save(path)
            first = extract_docx_ordered(str(path))
            second = extract_docx_ordered(str(path))

        first_units = first.units_as_dicts()
        second_units = second.units_as_dicts()
        formulas = [
            unit for unit in first_units
            if unit["unit_type"] == "formula"
            and unit["payload"].get("source_syntax") == "latex"
        ]
        citations = [unit for unit in first_units if unit["unit_type"] == "citation"]
        parent = next(unit for unit in first_units if unit["text"].startswith("公式"))
        self.assertEqual(len(formulas), 1)
        self.assertEqual(len(citations), 1)
        self.assertEqual(formulas[0]["relations"]["parent_unit_id"], parent["unit_id"])
        self.assertEqual(citations[0]["payload"]["numbers"], ["1", "2"])
        self.assertTrue(all(len(unit["content_hash"]) == 64 for unit in first_units))
        self.assertEqual(
            [(unit["unit_id"], unit["content_hash"]) for unit in first_units],
            [(unit["unit_id"], unit["content_hash"]) for unit in second_units],
        )
        self.assertEqual(first.metadata["source_sha256"], second.metadata["source_sha256"])


##### 成功与边界板块 #####


class FidelityPassTests(unittest.TestCase):
    def test_text_citation_formula_order_and_duplicate_text_pass(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            extract, recognize, render, _ = _virtual_case(Path(temp_dir))
            report = build_content_fidelity_report(extract, recognize, render)
        self.assertTrue(report["all_passed"])
        self.assertEqual(report["metrics"]["body_paragraph_coverage"]["status"], "pass")
        self.assertEqual(report["metrics"]["plain_text_integrity"]["status"], "pass")
        self.assertEqual(report["metrics"]["citation_coverage"]["status"], "pass")
        self.assertEqual(report["metrics"]["latex_formula_coverage"]["status"], "pass")
        self.assertEqual(report["metrics"]["formula_coverage"]["status"], "pass")
        self.assertEqual(report["metrics"]["formula_render_order"]["status"], "pass")

    def test_unverified_image_is_explicit_and_blocks(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            extract, recognize, render, _ = _virtual_case(Path(temp_dir), with_image=True)
            report = build_content_fidelity_report(extract, recognize, render)
        self.assertFalse(report["all_passed"])
        self.assertEqual(report["metrics"]["image_object_coverage"]["status"], "fail")

    def test_unresolved_citation_fails(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            extract, recognize, render, _ = _virtual_case(Path(temp_dir), unresolved=True)
            report = build_content_fidelity_report(extract, recognize, render)
        self.assertEqual(report["metrics"]["citation_coverage"]["status"], "fail")
        self.assertFalse(report["all_passed"])

    def test_transactional_literal_citation_fallback_is_degraded(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            extract, recognize, render, _ = _virtual_case(
                Path(temp_dir), unresolved=True,
            )
            render.render_trace["bibliography"] = {
                "transaction_status": "rolled_back",
                "all_citations_resolved": False,
                "required_numbers": ["1"],
                "body_citation_unit_ids": ["u-000003"],
            }
            report = build_content_fidelity_report(extract, recognize, render)
        self.assertEqual(
            report["metrics"]["citation_coverage"]["status"], "degraded",
        )
        self.assertEqual(
            report["metrics"]["bibliography_source_integrity"]["status"],
            "degraded",
        )
        self.assertTrue(report["all_passed"])


##### 故障注入板块 #####


class FidelityFailureTests(unittest.TestCase):
    def test_missing_formula_render_id_blocks_unified_formula_gate(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            extract, recognize, render, _ = _virtual_case(Path(temp_dir))
            body = next(
                record for record in render.render_trace["records"]
                if record["marker_unit_id"] == "u-000001"
            )
            body["details"]["rendered_formula_ids"] = []
            body["details"]["rendered_inline_formula_ids"] = []
            report = build_content_fidelity_report(extract, recognize, render)
        self.assertEqual(report["metrics"]["formula_coverage"]["status"], "fail")
        self.assertEqual(report["metrics"]["formula_render_order"]["status"], "fail")
        self.assertFalse(report["all_passed"])

    def test_missing_unit_marker_is_detected(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            extract, recognize, render, path = _virtual_case(Path(temp_dir))
            text = path.read_text(encoding="utf-8")
            text = re.sub(
                r"% SCHOLAR_UNIT_BEGIN u-000001 body_paragraph.*?% SCHOLAR_UNIT_END u-000001\n?",
                "",
                text,
                flags=re.DOTALL,
            )
            path.write_text(text, encoding="utf-8")
            report = build_content_fidelity_report(extract, recognize, render)
        self.assertEqual(report["metrics"]["body_paragraph_coverage"]["status"], "fail")

    def test_reordered_units_fail_even_when_counts_match(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            extract, recognize, render, path = _virtual_case(Path(temp_dir))
            text = path.read_text(encoding="utf-8")
            blocks = re.findall(
                r"% SCHOLAR_UNIT_BEGIN u-00000[56] body_paragraph.*?% SCHOLAR_UNIT_END u-00000[56]",
                text,
                flags=re.DOTALL,
            )
            self.assertEqual(len(blocks), 2)
            text = text.replace(blocks[0], "__FIRST__").replace(blocks[1], blocks[0]).replace("__FIRST__", blocks[1])
            path.write_text(text, encoding="utf-8")
            report = build_content_fidelity_report(extract, recognize, render)
        self.assertEqual(report["metrics"]["render_order"]["status"], "fail")

    def test_unexpected_text_and_duplicate_marker_fail(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            extract, recognize, render, path = _virtual_case(Path(temp_dir))
            original = path.read_text(encoding="utf-8")
            path.write_text(original + "\n无来源新增正文", encoding="utf-8")
            unexpected = build_content_fidelity_report(extract, recognize, render)
            block = re.search(
                r"% SCHOLAR_UNIT_BEGIN u-000005 body_paragraph.*?% SCHOLAR_UNIT_END u-000005",
                original,
                flags=re.DOTALL,
            ).group(0)
            path.write_text(original + "\n" + block, encoding="utf-8")
            duplicate = build_content_fidelity_report(extract, recognize, render)
        self.assertEqual(unexpected["metrics"]["unexpected_output"]["status"], "fail")
        self.assertEqual(duplicate["metrics"]["unexpected_output"]["status"], "fail")


if __name__ == "__main__":
    unittest.main(verbosity=2)
