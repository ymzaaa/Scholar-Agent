"""完整确认结构的层级、归属、版本与无副作用校验。"""

import copy

import pytest
from docx import Document

from pipeline_api import PipelineRunner
from pipeline.confirmed_structure import (
    ConfirmedStructureError, apply_confirmed_structure, build_structure_snapshot,
    finalize_structure_snapshot, validate_object_bindings,
)


##### 完整快照板块 #####


def prepare(tmp_path):
    document = Document()
    document.add_heading("第一章 测试", 1)
    document.add_heading("1.1 数据", 2)
    document.add_heading("1.1.1 来源", 3)
    document.add_paragraph("不可修改的正文。")
    document.add_heading("参考文献", 1)
    for number in range(1, 4):
        document.add_paragraph(f"[{number}] Author. Paper[J]. Journal, 2024.")
    path = tmp_path / "structure.docx"
    document.save(path)
    runner = PipelineRunner()
    extracted = runner.extract(str(path))
    recognized = runner.recognize(extracted)
    source = extracted.extraction_report["metadata"]["source_sha256"]
    snapshot = build_structure_snapshot(extracted, recognized, source)
    return extracted, recognized, snapshot, source


def test_confirmed_snapshot_contains_ranges_and_reloads_without_recognition(tmp_path, monkeypatch):
    extracted, recognized, snapshot, source = prepare(tmp_path)
    before = copy.deepcopy((extracted, recognized, snapshot))
    confirmed = finalize_structure_snapshot(snapshot)
    confirmed["revision"] = 1
    assert confirmed["chapters"][0]["range"] == [0, 4]
    assert confirmed["chapters"][0]["source_number"] == "1"
    assert confirmed["headings"][1]["title"] == "数据"
    monkeypatch.setattr("pipeline.structure_recognizer.infer_heading", lambda *a, **kw: pytest.fail("确认应用不能识别标题"))
    result = apply_confirmed_structure(extracted, None, confirmed, source_sha256=source, structure_revision=1)
    assert result.review["chapters"] == confirmed["chapters"]
    assert (extracted, recognized, snapshot) == before
    assert "source_sha256" not in confirmed
    assert result.internal == {"structure_revision": 1}


@pytest.mark.parametrize("levels,message", [
    (["body", "body", "body"], "至少需要一个章"),
    (["body", "section", "subsection"], "位于某一章"),
    (["chapter", "body", "subsection"], "首个节"),
])
def test_invalid_hierarchy_blocks_at_confirmation(tmp_path, levels, message):
    _, _, snapshot, _ = prepare(tmp_path)
    for heading, level in zip(snapshot["headings"], levels):
        heading.update(level=level, overridden=True, review_status="resolved")
    with pytest.raises(ConfirmedStructureError, match=message):
        finalize_structure_snapshot(snapshot)


@pytest.mark.parametrize("field,value,message", [
    ("project_source_sha256", "wrong", "哈希"),
    ("schema_version", "1.6.0", "版本"),
    ("revision", 2, "修订号"),
    ("chapters", [], "章节范围"),
])
def test_stale_or_incomplete_snapshot_is_rejected(tmp_path, field, value, message):
    extracted, recognized, snapshot, source = prepare(tmp_path)
    snapshot = finalize_structure_snapshot(snapshot)
    snapshot["revision"] = 1
    snapshot[field] = value
    with pytest.raises(ConfirmedStructureError, match=message):
        apply_confirmed_structure(extracted, recognized, snapshot, source_sha256=source, structure_revision=1)


##### 对象确认板块 #####


def test_unbound_duplicate_and_layout_table_are_rejected():
    index = {"c1": {"unit_type": "paragraph"}, "c2": {"unit_type": "paragraph"},
             "t": {"unit_type": "table", "semantic_role": "data"}}
    binding = {"caption_unit_id": "c1", "kind": "table", "status": "unbound", "object_unit_ids": []}
    with pytest.raises(ConfirmedStructureError, match="图表关系尚未确认"):
        validate_object_bindings([binding], index)
    binding.update(status="bound", object_unit_ids=["t"])
    with pytest.raises(ConfirmedStructureError, match="重复绑定"):
        validate_object_bindings([binding, {**binding, "caption_unit_id": "c2"}], index)
    index["t"]["semantic_role"] = "figure_layout"
    with pytest.raises(ConfirmedStructureError, match="只能绑定数据表"):
        validate_object_bindings([binding], index)
