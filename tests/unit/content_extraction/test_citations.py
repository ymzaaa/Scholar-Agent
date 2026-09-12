"""数字上标、引用分类与集中确认的确定性边界。"""

import copy
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from content_extraction.citations import apply_citation_overrides, detect_word_reference_list
from pipeline.bibliography import write_word_bibliography
from pipeline_api import PipelineRunner
from recognition.structure import StructureValidationError, commit_structure_confirmation


##### 引用抽取板块 #####


def citation_document():
    """提供上标引用、数学反例、歧义文本和带空行的文末编号。"""
    doc = Document()
    doc.add_heading("第一章 测试", 1)
    p = doc.add_paragraph("已有研究")
    p.add_run("1,2-3").font.superscript = True
    p.add_run("。")
    p = doc.add_paragraph("值为 x")
    p.add_run("2").font.superscript = True
    p.add_run("，正文数字 3。")
    p = doc.add_paragraph("带脚注")
    run = p.add_run()
    run.font.superscript = True
    note = OxmlElement("w:footnoteReference")
    note.set(qn("w:id"), "1")
    run._r.append(note)
    doc.add_paragraph("符号[1]位于此处")
    doc.add_paragraph("已有研究[2-3]。")
    for i in range(8):
        doc.add_paragraph(f"正文{i}。")
    doc.add_heading("参考文献", 1)
    for i in range(1, 4):
        doc.add_paragraph(f"[{i}] Author. Paper[J]. Journal, 2024.")
        doc.add_paragraph("")
    return doc


def test_numeric_superscripts_brackets_and_math_exclusions(tmp_path):
    path = tmp_path / "citations.docx"
    citation_document().save(path)
    runner = PipelineRunner()
    extracted = runner.extract(str(path))
    recognized = runner.recognize(extracted)
    units = [u for u in extracted.content_units if u["unit_type"] == "citation"]
    superscripts = [u for u in units if u["payload"]["citation_syntax"] == "numeric_superscript"]
    assert len(superscripts) == 1
    assert superscripts[0]["text"] == "1,2-3"
    assert superscripts[0]["payload"]["numbers"] == ["1", "2", "3"]
    assert recognized.review["citation_review"]["counts"] == {
        "body_citation": 2, "non_citation": 0, "reference_label": 3, "needs_review": 1,
    }
    assert len(detect_word_reference_list(extracted.paragraphs)["entries"]) == 3
    with pytest.raises(ValueError, match="引用候选未确认"):
        runner.render(extracted, recognized, "missing-template", str(tmp_path / "out"), None)
    assert not (tmp_path / "out").exists()


def test_superscript_without_binding_stays_reviewable(tmp_path):
    doc = Document()
    doc.add_paragraph("已有研究").add_run("8").font.superscript = True
    path = tmp_path / "unbound.docx"
    doc.save(path)
    runner = PipelineRunner()
    recognized = runner.recognize(runner.extract(str(path)))
    assert recognized.review["citation_review"]["counts"]["needs_review"] == 1


def test_bibliography_label_width_tracks_largest_number(tmp_path):
    entries = [{"number": str(i), "raw_text": f"条目{i}"} for i in range(1, 121)]
    plan = {"word_reference_list": {"entries": entries}, "number_to_key": {str(i): f"r{i}" for i in range(1, 121)}}
    path = write_word_bibliography(tmp_path, plan)
    content = (tmp_path / path).read_text(encoding="utf-8")
    assert content.startswith(r"\begin{thebibliography}{999}")
    assert content.count(r"\bibitem[") == 120
    assert r"\bibitem[120]{r120}" in content


##### 确认事务板块 #####


def _review():
    return {"decisions": [{"unit_id": "u-000001", "decision": "needs_review", "raw": "[1]", "numbers": ["1"], "source": {"start_offset": 1}}]}


@pytest.mark.parametrize("overrides", [
    [{"unit_id": "unknown", "decision": "body_citation"}],
    [{"unit_id": "u-000001", "decision": "reference_label"}],
    [{"unit_id": "u-000001", "decision": "body_citation", "raw": "[2]"}],
])
def test_invalid_override_is_rejected(overrides):
    with pytest.raises(ValueError):
        apply_citation_overrides(_review(), overrides)


def test_confirmation_preserves_original_evidence():
    original = _review()
    result = apply_citation_overrides(original, [{"unit_id": "u-000001", "decision": "non_citation"}])
    assert original["decisions"][0]["decision"] == "needs_review"
    assert result["decisions"][0]["raw"] == "[1]"
    assert result["decisions"][0]["numbers"] == ["1"]
    assert result["decisions"][0]["source"] == {"start_offset": 1}
    assert result["all_candidates_classified"]


def test_structure_transaction_citation_lock_and_unresolved_guard(tmp_path, monkeypatch):
    snapshot = {"revision": 0, "project_source_sha256": "hash", "confirmed": False,
                "schema_version": "1.7.0", "regions": {"body_range": [0, 1]},
                "headings": [{"unit_id": "u-000002", "text": "第一章 测试", "title": "测试", "level": "chapter", "confidence": 1.0, "evidence": [], "source": {"paragraph_index": 0}}],
                "object_bindings": [], "unit_index": {}, "citation_review": _review()}
    monkeypatch.setattr("recognition.structure.read_structure_snapshot", lambda _: copy.deepcopy(snapshot))
    writes = []
    monkeypatch.setattr("recognition.structure.write_confirmed_snapshot", lambda _, value: writes.append(value))
    project = SimpleNamespace(project_id="p", source_sha256="hash")
    store = SimpleNamespace(update_structure_state=lambda *args, **kwargs: None)
    kwargs = dict(project=project, store=store, base_revision=0, heading_overrides=[])
    with pytest.raises(StructureValidationError, match="引用候选未确认"):
        commit_structure_confirmation(**kwargs)
    assert not writes
    result = commit_structure_confirmation(**kwargs, citation_overrides=[{"unit_id": "u-000001", "decision": "body_citation"}])
    assert result["revision"] == 1 and result["confirmed"]
    snapshot.update(result)
    with pytest.raises(StructureValidationError, match="其他修订"):
        commit_structure_confirmation(**kwargs)
    project.source_sha256 = "changed"
    with pytest.raises(StructureValidationError, match="源文件不匹配"):
        commit_structure_confirmation(**{**kwargs, "base_revision": 1})
