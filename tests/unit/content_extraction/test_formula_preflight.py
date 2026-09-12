"""OMML 精确转换和复制模板前的公式阻断。"""

import pytest
from docx import Document
from docx.oxml import parse_xml

from content_extraction.formula_validation import FormulaPreflightBlocked, ensure_formula_preflight
from content_extraction.omml import OmmlParseError, omml_to_latex
from pipeline_api import PipelineRunner


##### OMML 转换板块 #####


def _math(body):
    return '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">' + body + '</m:oMath>'


@pytest.mark.parametrize("body, expected", [
    ('<m:f><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r></m:den></m:f>', r'\frac{a}{b}'),
    ('<m:f><m:fPr><m:type m:val="noBar"/></m:fPr><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r></m:den></m:f>', r'{a \atop b}'),
    ('<m:phant><m:e><m:r><m:t>x</m:t></m:r></m:e></m:phant>', r'\phantom{x}'),
    ('<m:borderBox><m:e><m:r><m:t>x</m:t></m:r></m:e></m:borderBox>', r'\boxed{x}'),
    ('<m:box><m:e><m:r><m:t>x</m:t></m:r></m:e></m:box>', 'x'),
])
def test_omml_known_structures(body, expected):
    assert omml_to_latex(_math(body)) == expected


@pytest.mark.parametrize("body", [
    '<m:unknown/>',
    '<m:bar><m:barPr><m:pos m:val="side"/></m:barPr><m:e><m:r><m:t>x</m:t></m:r></m:e></m:bar>',
    '<m:bar><m:barPr><m:pos m:val=""/></m:barPr><m:e><m:r><m:t>x</m:t></m:r></m:e></m:bar>',
    '<m:f><m:fPr><m:type m:val="skw"/></m:fPr><m:num/><m:den/></m:f>',
])
def test_unreliable_omml_fails(body):
    with pytest.raises(OmmlParseError):
        omml_to_latex(_math(body))


##### 渲染前阻断板块 #####


@pytest.mark.parametrize("review", [None, {}, {"status": "unknown"}, {"status": "blocked", "blocking": False}])
def test_only_explicit_pass_allows_render(review):
    with pytest.raises(FormulaPreflightBlocked):
        ensure_formula_preflight(review)
    ensure_formula_preflight({"status": "passed"})


@pytest.mark.parametrize("kind", ["omml", "latex", "missing"])
def test_formula_failure_blocks_before_template_access(tmp_path, kind):
    document = Document()
    document.add_heading("第一章 测试", 1)
    paragraph = document.add_paragraph("异常 $x" if kind == "latex" else "正文")
    if kind == "omml":
        paragraph._p.append(parse_xml(_math('<m:unknown/>')))
    path = tmp_path / "formula.docx"
    document.save(path)
    runner = PipelineRunner()
    extracted = runner.extract(str(path))
    recognized = runner.recognize(extracted)
    if kind == "missing":
        recognized.review.pop("formula_review")
    else:
        assert recognized.review["formula_review"]["malformed_formula_count"] == 1
    with pytest.raises(FormulaPreflightBlocked):
        runner.render(extracted, recognized, "missing-template", str(tmp_path / "output"), None)
    assert not (tmp_path / "output").exists()
