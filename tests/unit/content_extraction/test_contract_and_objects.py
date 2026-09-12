"""稳定抽取契约、块级控件及语义失败边界。"""

import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from content_extraction.persistence import load_extraction, save_extraction
from content_extraction.word_semantics import parse_field_instruction
from pipeline_api import PipelineRunner


##### 持久化契约板块 #####


def test_saved_contract_matches_schema_without_duplicate_fingerprints(tmp_path):
    document = Document()
    document.add_paragraph("内容 $x+1$。")
    path = tmp_path / "paper.docx"
    document.save(path)
    result = PipelineRunner().extract(str(path))
    digest = result.extraction_report["metadata"]["source_sha256"]
    target = save_extraction(result, tmp_path, source_sha256=digest)
    raw = (target / "content_units.json").read_text(encoding="utf-8")
    payload = json.loads(raw)
    schema = json.loads(Path("stage1/content_extraction/content_units.schema.json").read_text(encoding="utf-8"))
    assert set(payload) == set(schema["required"]) == {
        "schema_version", "source_path", "metadata", "units", "issues",
        "paragraphs", "tables", "media_files",
    }
    assert schema["additionalProperties"] is False
    assert schema["properties"]["schema_version"] == {"const": "1.7.0"}
    assert set(payload) == set(schema["properties"])
    for unit in payload["units"]:
        assert set(unit) == set(schema["$defs"]["content_unit"]["required"])
        assert unit["unit_type"] in schema["$defs"]["content_unit"]["properties"]["unit_type"]["enum"]
    assert payload["schema_version"] == "1.7.0"
    assert raw.count('"schema_version"') == raw.count('"source_sha256"') == 1
    assert set(payload["metadata"]) == set(result.extraction_report["metadata"])
    assert load_extraction(tmp_path, source_sha256=digest) == result


##### 块级遍历板块 #####


def _sdt(*nodes):
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    content.extend(nodes)
    control.append(content)
    return control


def test_nested_block_controls_keep_order_without_duplicate_cells(tmp_path):
    document = Document()
    first = document.add_paragraph("控件首段")._p
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "唯一单元格"
    last = document.add_paragraph("控件末段")._p
    document.element.body.insert(0, _sdt(first, _sdt(table._tbl), last))
    document.element.body.insert(1, OxmlElement("w:altChunk"))
    path = tmp_path / "controls.docx"
    document.save(path)
    result = PipelineRunner().extract(str(path))
    blocks = [u for u in result.content_units if u["unit_type"] in {"paragraph", "table"}]
    assert [u["unit_type"] for u in blocks] == ["paragraph", "table", "paragraph"]
    assert [p["text"] for p in result.paragraphs] == ["控件首段", "控件末段"]
    assert result.tables[0]["data"] == [["唯一单元格"]]
    issues = result.extraction_report["issues"]
    assert len(issues) == 1
    assert issues[0]["code"] == "BODY_BLOCK_UNSUPPORTED"
    assert "altChunk" in issues[0]["message"]


##### 语义降级板块 #####


def test_word_generated_contents_are_extracted_but_not_body_headings(tmp_path):
    document = Document()
    cached = document.add_heading("第一章 缓存目录标题", 1)._p
    control = _sdt(cached)
    properties = OxmlElement("w:sdtPr")
    part = OxmlElement("w:docPartObj")
    gallery = OxmlElement("w:docPartGallery")
    gallery.set(qn("w:val"), "Table of Contents")
    part.append(gallery)
    properties.append(part)
    control.insert(0, properties)
    document.element.body.insert(0, control)
    document.add_heading("第一章 正文标题", 1)
    document.add_paragraph("正文。")
    path = tmp_path / "contents.docx"
    document.save(path)
    runner = PipelineRunner()
    result = runner.extract(str(path))
    assert result.paragraphs[0]["text"] == "第一章 缓存目录标题"
    assert result.paragraphs[0]["generated_contents"] is True
    recognized = runner.recognize(result)
    assert [c["title"] for c in recognized.review["chapters"]] == ["正文标题"]
    assert len(recognized.review["heading_candidates"]) == 1


@pytest.mark.parametrize("instruction", ['HYPERLINK "file:///secret"', 'HYPERLINK "javascript:alert(1)"', 'DDE data', 'ADDIN data', 'DATE'])
def test_unsafe_links_and_unsupported_fields_degrade_during_extraction(tmp_path, instruction):
    document = Document()
    paragraph = document.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = "原始域结果"
    run.append(value)
    field.append(run)
    paragraph._p.append(field)
    path = tmp_path / "fields.docx"
    document.save(path)
    result = PipelineRunner().extract(str(path))
    unit = next(u for u in result.content_units if u["unit_type"] == "field")
    assert unit["status"] == "degraded"
    assert unit["text"] == "原始域结果"
    assert "dangerous" not in parse_field_instruction(instruction)
    assert result.extraction_report["issues"][0]["code"] == "FIELD_SEMANTIC_UNSUPPORTED"
